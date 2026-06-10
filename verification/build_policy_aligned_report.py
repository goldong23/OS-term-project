from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


def set_run_font(run, latin="Calibri", east_asia="Malgun Gothic", size=None):
    run.font.name = latin
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def set_style_font(style, latin="Calibri", east_asia="Malgun Gothic", size=None, color=None):
    style.font.name = latin
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "대표학번_가상메모리_페이지교체정책_보고서_최종개정.docx"
ASSETS = ROOT / "report_assets"
ASSETS.mkdir(exist_ok=True)

POLICIES = ["FIFO", "NUR (0,1 first)", "NUR (1,0 first)", "Second Chance", "LRFU-Lite"]
WORKLOADS = [
    ("Textbook", "123412512345", "강의자료 대표 예제이며 FIFO의 이상 현상과 추가 정책의 개선 효과를 확인하기 위한 입력이다."),
    ("NURPriority", "ABCDABEFABGHABCD", "modified bit 우선순위가 결과에 영향을 주는지 확인하기 위한 혼합 참조열이다."),
    ("ClockSensitive", "ABCDEABCDA", "Second Chance의 clock 시작 위치가 victim 선택에 영향을 주는지 확인하기 위한 입력이다."),
    ("FrequencyBias", "AAAABCAAADEFAAA", "A가 반복적으로 등장하여 LRFU-Lite의 빈도 보호 효과를 확인하기 위한 입력이다."),
    ("LocalityShift", "ABCABCABCDEFDEFABC", "ABC 작업 집합과 DEF 작업 집합이 전환될 때 정책별 적응성을 확인하기 위한 입력이다."),
]
FRAMES = [3, 4, 5]
DELAY_UNIT_MS = 10
LRFU_DECAY = 0.85


def policy_key(name):
    return {
        "FIFO": "FIFO",
        "NUR (0,1 first)": "NUR01",
        "NUR (1,0 first)": "NUR10",
        "Second Chance": "SC",
        "LRFU-Lite": "LRFU",
    }[name]


def simulate(reference, frames, policy, clock_start=1, reset_interval=4, modified_pages="AD"):
    frame = []
    fifo_queue = []
    r_bits = {}
    m_bits = {}
    scores = {}
    clock = (clock_start - 1) % frames
    hit = fault = migration = 0
    trace = []

    for idx, page in enumerate(reference):
        events = []
        if reset_interval > 0 and idx > 0 and idx % reset_interval == 0:
            for k in list(r_bits):
                r_bits[k] = False
            events.append(f"R reset before step {idx + 1}")

        if policy == "LRFU-Lite":
            for k in list(scores):
                scores[k] *= LRFU_DECAY

        victim = "-"
        status = ""
        if page in frame:
            hit += 1
            status = "Hit"
            r_bits[page] = True
            if page in modified_pages:
                m_bits[page] = True
            if policy == "LRFU-Lite":
                scores[page] = scores.get(page, 0.0) + 1.0
        else:
            fault += 1
            if len(frame) < frames:
                status = "Page Fault"
                frame.append(page)
                fifo_queue.append(page)
            else:
                status = "Migration"
                migration += 1
                if policy == "FIFO":
                    victim = fifo_queue.pop(0)
                    pos = frame.index(victim)
                    frame[pos] = page
                    fifo_queue.append(page)
                    events.append("FIFO selected oldest resident page")
                elif policy.startswith("NUR"):
                    order = [(0, 0), (0, 1), (1, 0), (1, 1)]
                    if policy == "NUR (1,0 first)":
                        order = [(0, 0), (1, 0), (0, 1), (1, 1)]
                    pos = 0
                    for rr, mm in order:
                        found = False
                        for i, candidate in enumerate(frame):
                            if int(r_bits.get(candidate, False)) == rr and int(m_bits.get(candidate, False)) == mm:
                                pos = i
                                found = True
                                break
                        if found:
                            break
                    victim = frame[pos]
                    frame[pos] = page
                    events.append(f"NUR class selected R={int(r_bits.get(victim, False))}, M={int(m_bits.get(victim, False))}")
                elif policy == "Second Chance":
                    scans = []
                    while True:
                        candidate = frame[clock]
                        scans.append(f"F{clock + 1}:{candidate}/R={int(r_bits.get(candidate, False))}")
                        if not r_bits.get(candidate, False):
                            pos = clock
                            victim = frame[pos]
                            frame[pos] = page
                            clock = (clock + 1) % frames
                            break
                        r_bits[candidate] = False
                        clock = (clock + 1) % frames
                    events.append("scan " + " -> ".join(scans))
                else:
                    pos = min(range(len(frame)), key=lambda i: scores.get(frame[i], 0.0))
                    victim = frame[pos]
                    events.append(f"minimum score {scores.get(victim, 0.0):.3f}")
                    frame[pos] = page
                for store in (r_bits, m_bits, scores):
                    store.pop(victim, None)

            r_bits[page] = True
            m_bits[page] = page in modified_pages
            if policy == "LRFU-Lite":
                scores[page] = 1.0

        snapshot = frame + ["-"] * (frames - len(frame))
        bit_state = " ".join(f"{p}(R={int(r_bits.get(p, False))},M={int(m_bits.get(p, False))})" for p in frame)
        score_state = " ".join(f"{p}={scores.get(p, 0.0):.2f}" for p in frame)
        trace.append((idx + 1, page, status, victim, " ".join(snapshot), bit_state if policy != "LRFU-Lite" else score_state, "; ".join(events)))

    return {
        "hit": hit,
        "fault": fault,
        "migration": migration,
        "fault_rate": round(fault / len(reference) * 100, 2),
        "delay": fault * DELAY_UNIT_MS,
        "trace": trace,
    }


def results():
    rows = []
    for name, ref, _ in WORKLOADS:
        for frame in FRAMES:
            for policy in POLICIES:
                r = simulate(ref, frame, policy)
                rows.append((name, ref, frame, policy, 1, 4, "AD", r["hit"], r["fault"], r["migration"], r["fault_rate"], r["delay"]))
    return rows


RESULTS = results()


def font(name="arial.ttf", size=22):
    for path in [Path("C:/Windows/Fonts") / name, Path("C:/Windows/Fonts/arial.ttf")]:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_bar(title, rows, path):
    width, height = 1100, 560
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    colors = ["#4E79A7", "#9C755F", "#B07AA1", "#59A14F", "#E15759"]
    title_font = font("arialbd.ttf", 27)
    label_font = font("arial.ttf", 16)
    d.text((80, 24), title, fill="#1F3A5F", font=title_font)
    x0, y0, w, h = 90, 475, 970, 350
    maxv = max(v for _, v in rows)
    d.line((x0, y0 - h, x0, y0), fill="#444444", width=2)
    d.line((x0, y0, x0 + w, y0), fill="#444444", width=2)
    for tick in range(0, maxv + 1, 2):
        y = y0 - int(h * tick / maxv)
        d.line((x0, y, x0 + w, y), fill="#E6EAF0")
        d.text((35, y - 9), str(tick), fill="#333333", font=label_font)
    gap = 30
    bw = int((w - gap * (len(rows) + 1)) / len(rows))
    for i, (name, value) in enumerate(rows):
        x = x0 + gap + i * (bw + gap)
        y = y0 - int(h * value / maxv)
        d.rectangle((x, y, x + bw, y0), fill=colors[i % len(colors)])
        d.text((x + bw // 2 - 8, y - 23), str(value), fill="#111111", font=label_font)
        label = name.replace("NUR ", "NUR\n").replace("Second Chance", "Second\nChance").replace("LRFU-Lite", "LRFU\nLite")
        d.multiline_text((x, y0 + 12), label, fill="#222222", font=font("arial.ttf", 13), spacing=2)
    img.save(path)


def draw_param(title, rows, path, x_label):
    width, height = 1000, 520
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = font("arialbd.ttf", 26)
    label_font = font("arial.ttf", 16)
    d.text((70, 24), title, fill="#1F3A5F", font=title_font)
    x0, y0, w, h = 90, 440, 850, 320
    maxv = max(v for _, v in rows)
    d.line((x0, y0 - h, x0, y0), fill="#444444", width=2)
    d.line((x0, y0, x0 + w, y0), fill="#444444", width=2)
    points = []
    for i, (name, value) in enumerate(rows):
        x = x0 + 60 + i * int((w - 120) / max(1, len(rows) - 1))
        y = y0 - int(h * value / maxv)
        points.append((x, y))
        d.ellipse((x - 6, y - 6, x + 6, y + 6), fill="#E15759")
        d.text((x - 8, y - 28), str(value), fill="#111111", font=label_font)
        d.text((x - 18, y0 + 12), str(name), fill="#222222", font=label_font)
    for a, b in zip(points, points[1:]):
        d.line((a[0], a[1], b[0], b[1]), fill="#E15759", width=3)
    d.text((width // 2 - 60, height - 35), x_label, fill="#333333", font=label_font)
    img.save(path)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_text(cell, text, bold=False, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    set_run_font(run, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, size=8.3, left_cols=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    left_cols = set(left_cols or [])
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, size)
        shade(t.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            cell_text(cells[i], value, False, size, align)
    return t


def para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, latin="Consolas", east_asia="Malgun Gothic", size=8.5)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        set_run_font(r)


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    for name, size, color in [("Normal", 10.5, "000000"), ("Heading 1", 15, "1F3A5F"), ("Heading 2", 12.5, "2E74B5"), ("Heading 3", 11.5, "1F4D78")]:
        s = doc.styles[name]
        set_style_font(s, size=size, color=color)


def title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("가상 메모리 페이지 교체기 설계 및 성능 분석")
    r.bold = True
    set_run_font(r, size=21)
    r.font.color.rgb = RGBColor.from_string("1F3A5F")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("소속: 컴퓨터공학과    학번: 제출 전 입력    이름: 제출 전 입력").bold = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("운영체제 Term Project 보고서").italic = True


def build():
    doc = Document()
    setup(doc)
    title(doc)
    heading(doc, "[요 약]", 1)
    para(doc, "본 프로젝트는 제공된 FIFO 페이지 교체 시뮬레이터를 기반으로 하여 과제에서 요구한 두 가지 추가 정책과 신규 정책을 구현하고, 각 정책의 page fault 발생 양상을 실험적으로 분석한 결과물이다. 최종 구현 정책은 FIFO, NUR (0,1 first), NUR (1,0 first), Second Chance, LRFU-Lite이며, 과제 요구 범위에 맞지 않는 기존 후보 정책은 최종 구현과 분석 대상에서 제외하였다.")
    para(doc, "NUR 정책은 reference bit와 modified bit를 사용한다. 본 프로젝트에서는 두 변형을 모두 구현하였다. 첫 번째 변형은 class 우선순위를 (0,0), (0,1), (1,0), (1,1) 순서로 두어 modified 되었지만 최근 사용되지 않은 page를 referenced clean page보다 먼저 교체한다. 두 번째 변형은 (0,0), (1,0), (0,1), (1,1) 순서로 두어 clean page를 더 우선적으로 교체한다. Second Chance는 clock hand 시작 frame을 UI에서 입력받도록 구현하였다. LRFU-Lite는 신규 정책으로, 최근성과 빈도성을 감쇠 점수로 결합한다.")
    para(doc, "보고서에서는 입력 요소를 서술형으로 설명하고, 결과 분석에서는 표와 그래프를 함께 제시하였다. 특히 NUR의 R reset interval, modified pages 입력, Second Chance의 clock start 같은 알고리즘 고유 매개변수를 변화시키며 결과 변화의 원인을 분석하였다. 또한 전체 소스 코드를 덤프하지 않고 핵심 알고리즘 이해에 필요한 코드 부분만 선별하여 제시하였다.")
    para(doc, "키워드: Virtual Memory, Page Replacement, FIFO, NUR, Not Used Recently, Second Chance, Clock Algorithm, LRFU-Lite, Reference Bit, Modified Bit")

    heading(doc, "I. 서론", 1)
    para(doc, "운영체제의 가상 메모리 시스템은 제한된 물리 메모리를 효율적으로 사용하기 위해 page 단위의 적재와 교체를 수행한다. 프로세스가 참조한 page가 물리 메모리에 없으면 page fault가 발생하고, 빈 frame이 없다면 기존 page 중 하나를 선택해 교체해야 한다. 어떤 page를 선택하느냐에 따라 이후 page fault 횟수와 실행 지연이 달라지므로 page replacement policy는 가상 메모리 성능을 결정하는 핵심 요소이다.")
    para(doc, "본 프로젝트의 목표는 FIFO만 제공된 시뮬레이터를 확장하여 수업에서 학습한 page replacement 정책을 직접 구현하고, 정책별 성능을 비교하는 것이다. 과제 요구사항에 맞추어 FIFO 외에 NUR 두 변형과 Second Chance를 구현하고, 신규 정책으로 LRFU-Lite를 유지하였다. UI도 정책별 매개변수를 받을 수 있도록 수정하였다.")
    para(doc, "구현의 핵심은 단순히 최종 page fault 수를 계산하는 데 있지 않다. 각 reference step에서 frame 내부 상태가 어떻게 바뀌는지, 어떤 bit와 pointer가 victim 선택에 영향을 주는지, 가변 인자 변경이 결과에 어떤 영향을 주는지를 보여주는 것이 중요하다. 따라서 프로그램은 step별 frame snapshot, R/M bit 상태, clock scan 결과, LRFU score를 출력하도록 확장하였다.")

    heading(doc, "II. 배경 지식 및 관련 기술", 1)
    para(doc, "Demand paging 환경에서 page는 실제로 참조될 때 메모리에 올라온다. 이 방식은 초기 적재 비용을 줄이고 물리 메모리 사용량을 절약하지만, 필요한 page가 메모리에 없을 때 page fault를 발생시킨다. Page fault는 일반 메모리 접근보다 훨씬 큰 비용을 가지므로 운영체제는 page fault 횟수를 줄이기 위한 교체 정책을 사용한다.")
    para(doc, "NUR, 즉 Not Used Recently 알고리즘은 page table entry에 존재할 수 있는 reference bit와 modified bit를 활용한다. Reference bit는 최근에 해당 page가 참조되었는지를 나타내고, modified bit는 page가 메모리에 올라온 뒤 수정되었는지를 나타낸다. 두 bit를 조합하면 page는 (0,0), (0,1), (1,0), (1,1) 네 class로 나뉜다. 일반적으로 (0,0)은 최근 사용되지 않았고 수정되지도 않았으므로 가장 교체하기 쉬운 class이다. (1,1)은 최근 사용되었고 수정되었으므로 가장 교체하기 부담스러운 class이다.")
    para(doc, "NUR에서 흥미로운 지점은 (0,1)과 (1,0)의 우선순위이다. (0,1)은 최근 사용되지 않았지만 modified 된 page이다. 최근 사용 여부만 보면 교체 후보에 가깝지만, dirty page라면 write-back 비용이 발생할 수 있다. 반대로 (1,0)은 최근 사용되었지만 수정되지 않은 clean page이다. 최근성 관점에서는 보호해야 하지만, 교체 비용 관점에서는 clean page라서 제거하기 쉽다. 따라서 어떤 우선순위를 선택하느냐에 따라 성능과 I/O 비용의 균형이 달라진다.")
    para(doc, "Second Chance는 FIFO의 단순한 queue 구조에 reference bit를 결합한 알고리즘이다. FIFO는 page의 최근 사용 여부를 고려하지 않기 때문에 오래된 page를 무조건 제거하지만, Second Chance는 reference bit가 1인 page를 만나면 bit를 0으로 내리고 한 번 더 기회를 준다. Clock 알고리즘은 이 과정을 원형 frame 배열과 clock hand로 구현한다. Clock hand가 어느 frame에서 시작하는지는 첫 victim 선택에 영향을 줄 수 있으므로, 본 프로젝트에서는 이를 UI 입력으로 받도록 하였다.")
    para(doc, "LRFU 계열은 최근성 정보와 빈도성 정보를 하나의 점수로 결합하려는 관련 연구 흐름에 속한다. 최근성 중심 정책은 방금 사용된 page를 잘 보호하지만 반복적으로 많이 사용되는 page를 장기적으로 보호하지 못할 수 있고, 단순 빈도 중심 정책은 오래된 참조 정보가 계속 남는 문제가 있다. LRFU-Lite는 이러한 아이디어를 학부 수준에서 구현하기 위해 단순화한 신규 정책이다. 각 page에 score를 두고 시간이 지날수록 score를 감쇠시키며, 참조가 발생하면 score를 증가시킨다.")

    heading(doc, "III. 본론: 주요 주제의 개요", 1)
    para(doc, "본 프로젝트의 주요 주제는 reference string을 입력받아 page replacement policy별 frame 변화와 성능 지표를 시뮬레이션하는 것이다. 사용자는 먼저 정책을 선택한다. 선택 가능한 정책은 FIFO, NUR (0,1 first), NUR (1,0 first), Second Chance, LRFU-Lite이다. 이후 reference string을 입력한다. Reference string은 한 글자를 하나의 page로 해석한다. 예를 들어 ABCDABEF는 A, B, C, D, A, B, E, F 순서의 page 참조를 의미한다.")
    para(doc, "Frame size는 물리 메모리에 동시에 유지할 수 있는 page 개수를 의미한다. Frame size가 작으면 같은 reference string에서도 교체가 자주 발생하고 page fault rate가 높아진다. 본 프로젝트는 frame size를 3, 4, 5로 바꾸어 결과 변화를 분석하였다. Second Chance 정책에서는 clock start를 별도로 입력한다. 이 값은 clock hand가 처음 검사할 frame 번호이며, 1부터 frame size까지의 양의 정수로 입력한다.")
    para(doc, "NUR 정책에는 R reset interval과 Modified Pages 입력이 추가된다. R reset interval은 몇 번의 참조마다 전체 reference bit를 0으로 초기화할지 결정한다. 이 값이 너무 작으면 최근 사용 정보가 빨리 사라지고, 너무 크면 거의 모든 page의 R bit가 1로 유지되어 class 구분력이 떨어진다. Modified Pages는 수정 연산이 발생한다고 가정할 page 문자 집합이다. 예를 들어 AD라고 입력하면 A와 D가 참조될 때 M bit가 1이 된다.")
    para(doc, "출력은 알고리즘 특성에 맞추어 다르게 보이도록 구성하였다. FIFO는 현재 FIFO order를 보여준다. NUR은 각 resident page의 R/M bit 상태를 A(R=1,M=0) 같은 형태로 보여준다. Second Chance는 현재 clock 위치와 scan 과정, 다음 clock 위치를 출력한다. LRFU-Lite는 page별 score를 출력한다. 공통적으로 hit count, page fault count, migration count, page fault rate, 실행 시간, page fault 지연 추정값을 출력한다.")

    heading(doc, "IV. 핵심 알고리즘 및 기능", 1)
    para(doc, "핵심 처리 흐름은 모든 정책에서 동일하다. 각 reference step마다 현재 page가 frame에 존재하는지 먼저 검사한다. 존재하면 hit로 처리하고 해당 정책의 메타데이터를 갱신한다. 존재하지 않으면 page fault가 발생한다. 빈 frame이 있으면 새 page를 삽입하고, 빈 frame이 없으면 정책별 victim selection 알고리즘으로 교체할 page를 선택한다.")
    code(doc, "for each page in referenceString:\n    if page is in frame:\n        hit++\n        update policy metadata\n    else:\n        fault++\n        if frame has empty slot:\n            insert page\n        else:\n            victim = SelectVictim(policy)\n            replace victim\n            migration++\n    record frame snapshot and algorithm state")
    para(doc, "FIFO는 제공된 알고리즘을 유지한 기준 정책이다. 교체가 필요하면 frame에 가장 먼저 들어온 page를 victim으로 선택한다. Hit가 발생해도 순서를 바꾸지 않는다. 따라서 FIFO는 구현 비용이 낮지만 reference bit나 modified bit를 활용하지 않기 때문에 workload에 따라 불리해질 수 있다.")
    code(doc, "FIFO victim selection:\n    victim = frame_window[0]\n    remove victim from front\n    append new page to rear")
    para(doc, "NUR (0,1 first)는 class 우선순위를 (0,0), (0,1), (1,0), (1,1) 순서로 둔다. 즉 최근 사용되지 않은 page를 우선하고, 그 안에서는 modified 되었더라도 R=0인 page를 R=1인 clean page보다 먼저 교체한다. 이 변형은 최근성 정보를 더 강하게 반영한다.")
    code(doc, "NUR (0,1 first):\n    priority = (0,0) -> (0,1) -> (1,0) -> (1,1)\n    scan frame in order\n    select first page whose (R,M) matches earliest class")
    para(doc, "NUR (1,0 first)는 class 우선순위를 (0,0), (1,0), (0,1), (1,1) 순서로 둔다. 이 변형은 clean page를 우선 교체하여 modified page의 write-back 가능성을 줄이는 관점에 가깝다. 따라서 동일한 reference string에서도 Modified Pages 입력에 따라 두 NUR 변형의 결과가 달라질 수 있다.")
    code(doc, "NUR (1,0 first):\n    priority = (0,0) -> (1,0) -> (0,1) -> (1,1)\n    scan frame in order\n    select first page whose (R,M) matches earliest class")
    para(doc, "Second Chance는 clock hand가 가리키는 frame부터 검사한다. R bit가 0이면 해당 page를 victim으로 선택한다. R bit가 1이면 0으로 내리고 다음 frame으로 이동한다. 이 과정은 victim이 선택될 때까지 반복된다. 본 프로젝트에서는 clock start를 UI에서 입력받아 clock hand의 초기 위치를 지정할 수 있게 했다.")
    code(doc, "Second Chance:\n    while true:\n        if R[clockHand] == 0:\n            victim = frame[clockHand]\n            clockHand = next frame\n            break\n        else:\n            R[clockHand] = 0\n            clockHand = next frame")
    para(doc, "LRFU-Lite는 신규 정책이다. 매 참조마다 resident page의 score를 0.85배로 감쇠시킨다. 참조 page가 hit되면 해당 score에 1.0을 더한다. Page fault가 발생하고 frame이 가득 차 있으면 score가 가장 낮은 page를 victim으로 선택한다. 이 정책은 최근성과 빈도성을 동시에 고려하려는 단순화된 recency-frequency 정책이다.")
    code(doc, "LRFU-Lite:\n    for each resident page p:\n        score[p] = score[p] * 0.85\n    if hit:\n        score[page] += 1.0\n    if replacement is needed:\n        victim = page with minimum score")

    heading(doc, "V. 성능 평가", 1)
    para(doc, "성능 평가는 Textbook, NURPriority, ClockSensitive, FrequencyBias, LocalityShift 다섯 가지 reference string으로 수행하였다. 입력 요소는 표로 나열하지 않고 여기서 서술한다. Textbook 입력은 강의자료의 대표 예제를 사용해 FIFO와 추가 정책의 차이를 확인하기 위한 것이다. NURPriority 입력은 R/M bit class 우선순위가 결과에 미치는 영향을 확인하기 위한 혼합 참조열이다. ClockSensitive 입력은 Second Chance의 clock start 변화가 결과에 영향을 주는지 확인하기 위한 입력이다. FrequencyBias 입력은 A가 반복적으로 나타나 LRFU-Lite의 빈도 보호 효과를 확인하기 위한 입력이다. LocalityShift 입력은 작업 집합이 ABC에서 DEF로 바뀌었다가 다시 돌아오는 상황을 표현한다.")
    para(doc, "기본 실험에서는 clock start를 1, R reset interval을 4, Modified Pages를 AD로 두었다. 이후 가변 인자 분석에서 clock start를 1,2,3,4로 변경하고, R reset interval을 2,3,4,6으로 변경하며, Modified Pages를 A, AD, BDF, empty로 변경하였다. 평가 기준은 hit count, page fault count, migration count, page fault rate, estimated delay이다.")
    table(doc, ["Workload", "Frame", "Policy", "Hit", "Fault", "Migration", "Fault Rate", "Delay"], [(w, f, p, h, fault, m, f"{rate:.2f}%", delay) for w, ref, f, p, c, rst, mod, h, fault, m, rate, delay in RESULTS if f == 4], size=7.6, left_cols={0, 2})

    for workload in ["Textbook", "NURPriority", "ClockSensitive", "FrequencyBias", "LocalityShift"]:
        path = ASSETS / f"aligned_{workload}.png"
        rows = [(p, fault) for w, ref, f, p, c, rst, mod, h, fault, m, rate, delay in RESULTS if w == workload and f == 4]
        draw_bar(f"{workload} workload, frame=4 page faults", rows, path)
        doc.add_picture(str(path), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading(doc, "VI. 결과 분석", 1)
    para(doc, "Textbook workload에서 frame=4 기준 FIFO는 fault 10회를 보였다. 반면 NUR 두 변형은 fault 6회, Second Chance는 fault 7회, LRFU-Lite는 fault 8회를 보였다. 이는 R bit를 사용하는 정책들이 단순 FIFO보다 최근 참조 정보를 더 잘 활용했기 때문이다. 특히 NUR은 R reset interval이 4로 설정되어 일정 주기마다 오래된 참조 정보를 제거하므로, 최근 사용 여부를 구분하면서도 class 기반 선택을 수행할 수 있었다.")
    para(doc, "NURPriority workload에서는 NUR (0,1 first)가 frame=4에서 fault 10회, NUR (1,0 first)가 fault 11회를 보였다. Modified Pages를 AD로 둔 조건에서는 (0,1)을 (1,0)보다 우선한 정책이 더 좋은 결과를 냈다. 이는 해당 reference string에서 modified 되었지만 최근 사용되지 않은 page를 먼저 제거하는 편이, 최근 사용된 clean page를 제거하는 것보다 이후 재참조 손실이 적었기 때문이다.")
    para(doc, "ClockSensitive workload에서는 Second Chance의 clock start가 결과에 영향을 주었다. clock start가 1 또는 2일 때 fault는 9회였지만, 3일 때 8회, 4일 때 7회로 감소했다. 이는 clock hand의 초기 위치가 첫 replacement 시점의 scan 순서를 바꾸고, 그 결과 이후 frame 구성까지 연쇄적으로 달라지기 때문이다. Second Chance는 FIFO처럼 순서 기반 요소가 있으면서도 R bit를 수정하므로 초기 pointer 위치가 작지만 분명한 영향을 줄 수 있다.")
    para(doc, "FrequencyBias workload에서는 A가 반복적으로 등장한다. frame=4 기준 FIFO는 fault 7회였고 NUR, Second Chance, LRFU-Lite는 모두 fault 6회를 보였다. 반복 참조가 강한 상황에서는 R bit 기반 정책과 LRFU-Lite 모두 자주 사용되는 page를 보호하는 효과를 낸다. LRFU-Lite는 A의 score가 반복적으로 증가하므로 A를 victim으로 선택하지 않는 경향이 생긴다.")
    para(doc, "LocalityShift workload에서는 작업 집합이 ABC에서 DEF로 이동했다가 다시 ABC로 돌아온다. 이 상황에서 frame=4 기준 FIFO와 Second Chance는 fault 9회, NUR (0,1 first)는 fault 9회, NUR (1,0 first)는 fault 10회, LRFU-Lite는 fault 12회를 보였다. LRFU-Lite는 과거에 자주 참조된 page의 score를 어느 정도 유지하기 때문에 작업 집합이 급격히 바뀌는 상황에서는 민첩성이 떨어질 수 있다. 이는 빈도 기반 정책의 전형적인 한계와 연결된다.")

    heading(doc, "VII. 가변 인자 변경에 따른 상관 분석", 1)
    clock_rows = []
    for c in [1, 2, 3, 4]:
        r = simulate("ABCDEABCDA", 4, "Second Chance", c, 4, "AD")
        clock_rows.append((c, r["fault"]))
    path = ASSETS / "aligned_clock_param.png"
    draw_param("Second Chance clock start sensitivity", clock_rows, path, "Clock Start")
    doc.add_picture(str(path), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Clock start가 커질수록 이 실험에서는 fault가 감소하는 추세를 보였다. 이는 특정 reference string에서 뒤쪽 frame을 먼저 검사하는 것이 첫 replacement 이후 더 유리한 frame 구성을 만들었기 때문이다. 다만 이는 모든 입력에서 보장되는 일반 법칙은 아니다. Clock start는 reference string과 현재 frame 구성에 따라 victim 순서를 바꾸는 매개변수이므로, 결과와의 관계는 입력 의존적이다.")

    reset_rows = []
    for rst in [2, 3, 4, 6]:
        r = simulate("ABCDABEFABGHABCD", 4, "NUR (0,1 first)", 1, rst, "AD")
        reset_rows.append((rst, r["fault"]))
    path = ASSETS / "aligned_reset_param.png"
    draw_param("NUR R reset interval sensitivity", reset_rows, path, "R Reset Interval")
    doc.add_picture(str(path), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "R reset interval은 NUR의 class 구분력과 직접 연결된다. reset interval이 2 또는 3일 때는 fault 11회였고, 4일 때 fault 10회로 개선되었지만, 6에서는 다시 11회가 되었다. 너무 자주 reset하면 최근 참조 정보가 빨리 사라지고, 너무 늦게 reset하면 대부분의 page가 R=1에 머물러 class 차이가 약해진다. 따라서 중간 정도의 reset 주기가 더 좋은 결과를 낼 수 있다.")

    mod_rows = []
    for mod in ["A", "AD", "BDF", ""]:
        r = simulate("ABCDABEFABGHABCD", 4, "NUR (1,0 first)", 1, 4, mod)
        mod_rows.append((mod if mod else "none", r["fault"]))
    path = ASSETS / "aligned_modified_param.png"
    draw_param("NUR modified pages sensitivity", mod_rows, path, "Modified Pages")
    doc.add_picture(str(path), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Modified Pages 입력은 NUR의 M bit class를 바꾸므로 victim 선택에 영향을 준다. NUR (1,0 first)는 clean referenced page를 (0,1)보다 먼저 볼 수 있기 때문에, 어떤 page가 modified로 표시되는지에 따라 결과가 달라졌다. AD 또는 BDF를 modified로 둔 경우 fault가 11회였고, A만 modified이거나 modified page가 없는 경우 fault가 10회였다. 이는 M bit가 단순 보조 정보가 아니라 class 우선순위와 결합되어 실제 교체 순서를 바꾸는 입력임을 보여준다.")

    heading(doc, "VIII. 핵심 코드 발췌", 1)
    para(doc, "보고서에는 전체 소스 코드를 모두 덤프하지 않고 핵심 알고리즘 이해에 필요한 부분만 발췌한다. 전체 구현은 제출 ZIP과 GitHub 저장소의 Core.cs, Form1.cs에서 확인할 수 있다.")
    code(doc, "private int SelectNurVictim(bool preferReferencedClean, Page historyPage)\n{\n    int[][] order = preferReferencedClean\n        ? new int[][] { (0,0), (1,0), (0,1), (1,1) }\n        : new int[][] { (0,0), (0,1), (1,0), (1,1) };\n    foreach class in order:\n        scan resident frames and select first matching (R,M) page;\n}")
    code(doc, "private int SelectSecondChanceVictim(Page historyPage)\n{\n    while (true) {\n        if (referenceBits[clockHandPage] == false) select it;\n        referenceBits[clockHandPage] = false;\n        clockHand = next frame;\n    }\n}")
    code(doc, "private int SelectLrfuLiteVictim(Page historyPage)\n{\n    victim = page with minimum lrfuScores[page];\n}\n\nprivate void DecayLrfuScores()\n{\n    foreach resident page: score *= 0.85;\n}")

    heading(doc, "IX. 결론", 1)
    para(doc, "본 프로젝트는 과제 요구사항에 맞추어 FIFO 외에 NUR 두 변형과 Second Chance를 구현하고, 신규 정책으로 LRFU-Lite를 유지하였다. 최종 프로그램은 Second Chance의 clock start, NUR의 R reset interval과 Modified Pages를 입력받으며, 각 알고리즘의 내부 상태를 출력한다.")
    para(doc, "실험 결과 단순 FIFO보다 R bit를 활용하는 정책들이 여러 workload에서 더 좋은 결과를 보였다. 그러나 NUR의 두 변형은 modified bit 우선순위에 따라 결과가 달라졌고, Second Chance는 clock start에 따라 결과가 변했다. LRFU-Lite는 빈도 편향 workload에서는 효과적이었지만 작업 집합이 빠르게 전환되는 workload에서는 불리할 수 있었다. 따라서 page replacement 정책의 성능은 알고리즘 자체뿐 아니라 입력 참조열의 성격과 매개변수 설정에 크게 의존한다.")
    heading(doc, "참고 자료", 1)
    table(doc, ["번호", "자료"], [
        ("[1]", "운영체제 강의자료 Ch3. Memory Management and Virtual Memory."),
        ("[2]", "Term Project - Page Replacement Policy Design.pdf."),
        ("[3]", "Page replacement algorithm, Wikipedia, https://en.wikipedia.org/wiki/Page_replacement_algorithm"),
        ("[4]", "Cache replacement policies, Wikipedia, https://en.wikipedia.org/wiki/Cache_replacement_policies"),
        ("[5]", "D. Lee et al., LRFU page replacement policy research, IEEE Transactions on Computers, 2001."),
    ], left_cols={1})

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
