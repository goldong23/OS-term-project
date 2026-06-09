from pathlib import Path
from collections import defaultdict

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "대표학번_가상메모리_페이지교체정책_보고서_차별화개정.docx"
ASSETS = ROOT / "report_assets"
ASSETS.mkdir(exist_ok=True)

POLICIES = ["FIFO", "Optimal", "LRU", "Second Chance", "LRFU-Lite"]
WORKLOADS = [
    ("Textbook", "123412512345", "강의자료에 제시된 대표 참조열. Belady's anomaly와 Optimal/LRU 비교에 적합하다."),
    ("Locality", "ABCABCABCDEFDEFABC", "A/B/C와 D/E/F 작업 집합이 교대로 나타나는 시간 지역성 패턴이다."),
    ("Sequential Scan", "ABCDEFGHIJKL", "모든 페이지가 한 번씩만 참조되는 순차 스캔 패턴이다."),
    ("Mixed", "ABCDABEFABGHABCD", "A/B 재참조와 새로운 페이지 유입이 섞인 혼합 패턴이다."),
    ("Frequency Bias", "AAAABCAAADEFAAA", "A가 반복적으로 등장하는 빈도 편향 패턴으로 LRFU-Lite의 의도를 확인한다."),
]
FRAMES = [3, 4, 5]
DELAY_UNIT_MS = 10
LRFU_DECAY = 0.85


def simulate(reference, frames, policy):
    frame = []
    fifo_queue = []
    last_used = {}
    freq = {}
    ref_bits = {}
    scores = {}
    clock = 0
    hit = fault = migration = 0
    rows = []

    for i, page in enumerate(reference):
        if policy == "LRFU-Lite":
            for k in list(scores):
                scores[k] *= LRFU_DECAY

        status = ""
        victim = "-"

        if page in frame:
            hit += 1
            status = "Hit"
            loc = frame.index(page) + 1
            last_used[page] = i
            freq[page] = freq.get(page, 0) + 1
            if page in ref_bits:
                ref_bits[page] = True
            if policy == "LRFU-Lite":
                scores[page] = scores.get(page, 0.0) + 1.0
        else:
            fault += 1
            if len(frame) < frames:
                status = "Page Fault"
                frame.append(page)
                fifo_queue.append(page)
                loc = len(frame)
            else:
                status = "Migration"
                migration += 1
                if policy == "FIFO":
                    old = fifo_queue.pop(0)
                    idx = frame.index(old)
                    victim = old
                    frame[idx] = page
                    fifo_queue.append(page)
                    loc = idx + 1
                elif policy == "Optimal":
                    farthest = -1
                    idx = 0
                    for pos, candidate in enumerate(frame):
                        try:
                            next_use = reference.index(candidate, i + 1)
                        except ValueError:
                            next_use = 10**9
                        if next_use > farthest:
                            farthest = next_use
                            idx = pos
                    victim = frame[idx]
                    frame[idx] = page
                    loc = idx + 1
                elif policy == "LRU":
                    idx = min(range(len(frame)), key=lambda p: last_used.get(frame[p], -1))
                    victim = frame[idx]
                    frame[idx] = page
                    loc = idx + 1
                elif policy == "Second Chance":
                    while True:
                        candidate = frame[clock]
                        if not ref_bits.get(candidate, False):
                            idx = clock
                            victim = frame[idx]
                            frame[idx] = page
                            loc = idx + 1
                            clock = (idx + 1) % frames
                            break
                        ref_bits[candidate] = False
                        clock = (clock + 1) % frames
                else:
                    idx = min(range(len(frame)), key=lambda p: scores.get(frame[p], 0.0))
                    victim = frame[idx]
                    frame[idx] = page
                    loc = idx + 1

                for store in (last_used, freq, ref_bits, scores):
                    store.pop(victim, None)

            last_used[page] = i
            freq[page] = 1
            ref_bits[page] = False
            if policy == "LRFU-Lite":
                scores[page] = 1.0

        snapshot = frame + ["-"] * (frames - len(frame))
        score_text = ""
        if policy == "LRFU-Lite":
            score_text = ", ".join(f"{p}:{scores.get(p, 0):.2f}" for p in frame)

        rows.append({
            "step": i + 1,
            "ref": page,
            "status": status,
            "victim": victim,
            "loc": loc,
            "snapshot": " ".join(snapshot),
            "score": score_text,
        })

    return {
        "hit": hit,
        "fault": fault,
        "migration": migration,
        "fault_rate": round(fault / len(reference) * 100, 2),
        "delay": fault * DELAY_UNIT_MS,
        "trace": rows,
    }


def all_results():
    rows = []
    for w_name, ref, _ in WORKLOADS:
        for frame_size in FRAMES:
            for policy in POLICIES:
                r = simulate(ref, frame_size, policy)
                rows.append((w_name, ref, frame_size, policy, r["hit"], r["fault"], r["migration"], r["fault_rate"], r["delay"]))
    return rows


RESULTS = all_results()


def font(name="arial.ttf", size=22):
    for path in [Path("C:/Windows/Fonts") / name, Path("C:/Windows/Fonts/arial.ttf")]:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_bar_chart(title, rows, path):
    width, height = 1100, 560
    margin_l, margin_r, margin_t, margin_b = 90, 40, 80, 105
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = font("arialbd.ttf", 27)
    label_font = font("arial.ttf", 18)
    small_font = font("arial.ttf", 15)
    colors = ["#4E79A7", "#59A14F", "#F28E2B", "#B07AA1", "#E15759"]
    d.text((margin_l, 24), title, fill="#1F3A5F", font=title_font)
    chart_w = width - margin_l - margin_r
    chart_h = height - margin_t - margin_b
    x0, y0 = margin_l, height - margin_b
    d.line((x0, margin_t, x0, y0), fill="#404040", width=2)
    d.line((x0, y0, width - margin_r, y0), fill="#404040", width=2)
    max_fault = max(row[1] for row in rows)
    step = 2 if max_fault <= 16 else 5
    for tick in range(0, max_fault + 1, step):
        y = y0 - int(chart_h * tick / max_fault)
        d.text((25, y - 10), str(tick), fill="#333333", font=small_font)
        d.line((x0, y, width - margin_r, y), fill="#E8EEF5", width=1)
    bar_gap = 30
    bar_w = int((chart_w - bar_gap * (len(rows) + 1)) / len(rows))
    for idx, (name, faults) in enumerate(rows):
        x = x0 + bar_gap + idx * (bar_w + bar_gap)
        y = y0 - int(chart_h * faults / max_fault)
        d.rectangle((x, y, x + bar_w, y0), fill=colors[idx % len(colors)])
        d.text((x + bar_w / 2 - 8, y - 24), str(faults), fill="#111111", font=label_font)
        label = name.replace("Second Chance", "Second\nChance").replace("LRFU-Lite", "LRFU\nLite")
        d.multiline_text((x, y0 + 12), label, fill="#222222", font=small_font, spacing=2)
    d.text((20, margin_t + 2), "Faults", fill="#333333", font=small_font)
    img.save(path)


def draw_frame_chart(workload, path):
    width, height = 1200, 620
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    title_font = font("arialbd.ttf", 27)
    label_font = font("arial.ttf", 16)
    colors = ["#4E79A7", "#59A14F", "#F28E2B", "#B07AA1", "#E15759"]
    d.text((70, 24), f"{workload} workload: frame size sensitivity", fill="#1F3A5F", font=title_font)
    x0, y0, chart_w, chart_h = 90, 500, 1030, 380
    d.line((x0, y0 - chart_h, x0, y0), fill="#404040", width=2)
    d.line((x0, y0, x0 + chart_w, y0), fill="#404040", width=2)
    subset = [r for r in RESULTS if r[0] == workload]
    max_fault = max(r[5] for r in subset)
    for tick in range(0, max_fault + 1, 2):
        y = y0 - int(chart_h * tick / max_fault)
        d.text((35, y - 10), str(tick), fill="#333333", font=label_font)
        d.line((x0, y, x0 + chart_w, y), fill="#E8EEF5", width=1)
    group_w = chart_w // len(FRAMES)
    bar_w = 34
    for f_idx, frame_size in enumerate(FRAMES):
        gx = x0 + f_idx * group_w + 40
        d.text((gx + 80, y0 + 18), f"Frame {frame_size}", fill="#222222", font=label_font)
        for p_idx, policy in enumerate(POLICIES):
            fault = next(r[5] for r in subset if r[2] == frame_size and r[3] == policy)
            x = gx + p_idx * (bar_w + 12)
            y = y0 - int(chart_h * fault / max_fault)
            d.rectangle((x, y, x + bar_w, y0), fill=colors[p_idx])
            d.text((x + 6, y - 20), str(fault), fill="#111111", font=font("arial.ttf", 13))
    legend_x = 825
    for idx, policy in enumerate(POLICIES):
        d.rectangle((legend_x, 105 + idx * 28, legend_x + 18, 123 + idx * 28), fill=colors[idx])
        d.text((legend_x + 28, 102 + idx * 28), policy, fill="#222222", font=label_font)
    img.save(path)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, font_size=8.5, left_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=font_size)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    left_cols = set(left_cols or [])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.LEFT if i in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[i], value, size=font_size, align=align)
    return table


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return p


def setup_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    for style_name, size, color in [
        ("Normal", 10.5, "000000"),
        ("Heading 1", 15, "1F3A5F"),
        ("Heading 2", 12.5, "2E74B5"),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        if style_name != "Normal":
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(5)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("가상 메모리 페이지 교체기 설계 및 성능 분석")
    run.bold = True
    run.font.size = Pt(21)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor.from_string("1F3A5F")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("소속: 컴퓨터공학과    학번: 제출 전 입력    이름: 제출 전 입력").bold = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("운영체제 Term Project 보고서").italic = True


def add_summary(doc):
    heading(doc, "[요 약]", 1)
    add_para(doc, "본 프로젝트에서는 운영체제 가상 메모리 단원에서 핵심적으로 다루는 페이지 교체 정책을 직접 구현하고, 동일한 reference string에서 정책별 page fault 발생 양상과 성능 차이를 분석하였다. 제공된 C# WinForms 기반 FIFO 시뮬레이터를 확장하여 FIFO, Optimal, LRU, Second Chance, LRFU-Lite 총 다섯 가지 정책을 선택 실행할 수 있도록 구성하였다. 프로그램은 사용자가 페이지 교체 정책, reference string, frame size를 입력하면 단계별 프레임 상태, Hit Count, Page Fault Count, Migration Count, Page Fault Rate, 실행 시간, Page Fault에 따른 지연 시간 추정값을 출력한다.")
    add_para(doc, "본 보고서는 단순히 구현 결과를 나열하는 데 그치지 않고, 각 알고리즘이 어떤 가정에서 출발하는지, 어떤 메타데이터를 필요로 하는지, 어떤 workload에서 성능이 좋아지거나 나빠지는지까지 분석한다. 특히 신규 정책 LRFU-Lite는 LRU와 LFU를 결합하려는 LRFU 계열의 아이디어를 학부 수준에서 구현 가능한 형태로 단순화한 정책이다. 이 정책은 resident page마다 점수를 두고 매 참조 시 기존 점수를 0.85배로 감쇠시킨 뒤 참조된 page에 1.0을 더한다. 교체 시에는 감쇠된 점수가 가장 낮은 page를 victim으로 선택한다.")
    add_para(doc, "실험은 강의자료 예제, 시간 지역성 패턴, 순차 스캔 패턴, 혼합 패턴, 빈도 편향 패턴 총 다섯 가지 reference string과 frame size 3, 4, 5를 사용하여 수행하였다. Textbook workload에서 frame size가 3에서 4로 증가할 때 FIFO의 page fault가 9회에서 10회로 증가하여 Belady's anomaly를 확인하였다. Optimal은 대부분의 workload에서 이론적 최저 fault 수를 보였고, LRU와 Second Chance는 temporal locality가 있는 입력에서 안정적인 성능을 보였다. LRFU-Lite는 빈도 편향 workload에서 FIFO보다 적은 fault를 보였으나, 작업 집합이 급격히 바뀌는 Locality workload에서는 감쇠 계수 선택에 따라 성능이 저하될 수 있음을 확인하였다.")
    add_para(doc, "키워드: Virtual Memory, Demand Paging, Page Replacement, FIFO, Optimal, LRU, Second Chance, LRFU, Page Fault, Belady's Anomaly")


def add_intro(doc):
    heading(doc, "I. 서론", 1)
    heading(doc, "1. 프로젝트 배경과 개요", 2)
    for text in [
        "운영체제는 제한된 물리 메모리를 여러 프로세스가 효율적으로 사용할 수 있도록 가상 메모리 체계를 제공한다. 가상 메모리에서는 프로세스가 사용하는 주소 공간을 page 단위로 나누고, 실제 물리 메모리에는 필요한 page만 적재한다. 이 구조 덕분에 프로세스는 실제 물리 메모리보다 큰 주소 공간을 사용할 수 있지만, 필요한 page가 물리 메모리에 없을 때 page fault가 발생한다.",
        "Page fault가 발생했을 때 빈 frame이 존재하면 새 page를 적재하면 된다. 그러나 모든 frame이 사용 중이라면 기존 page 중 하나를 선택하여 내보내야 한다. 이 선택 규칙이 page replacement policy이다. 교체 정책의 선택은 전체 시스템 성능에 큰 영향을 준다. 잘못된 page를 내보내면 곧바로 다시 page fault가 발생하고, 디스크 I/O와 context switch 비용이 반복된다.",
        "본 프로젝트는 페이지 교체 정책을 이론으로만 이해하는 것이 아니라 실제 시뮬레이터로 구현하고, 동일한 reference string에서 정책별 결과가 어떻게 달라지는지 관찰하기 위해 수행되었다. 단순한 실행 결과보다 중요한 것은 결과의 원인을 설명하는 것이다. 따라서 본 보고서에서는 각 정책의 핵심 아이디어, 구현 방식, 단계별 동작 사례, 성능 평가 결과, 결과 변화 원인까지 함께 다룬다.",
    ]:
        add_para(doc, text)
    heading(doc, "2. 프로젝트 목표 및 의의", 2)
    for item in [
        "수업시간에 학습한 FIFO, Optimal, LRU, Second Chance 정책을 C# 프로그램으로 구현한다.",
        "Reference string과 frame size를 바꾸면서 Hit, Page Fault, Migration, Fault Rate 변화를 측정한다.",
        "단계별 frame snapshot을 출력하여 알고리즘의 동작 순서를 시각적으로 확인한다.",
        "그래프와 표를 통해 정책별 성능을 정량적으로 비교하고, 결과가 발생한 원인을 분석한다.",
        "수업에서 배운 정책을 확장하여 새로운 정책 LRFU-Lite를 제안하고 장단점을 평가한다.",
    ]:
        add_bullet(doc, item)
    heading(doc, "3. 구현하고자 하는 핵심 기능", 2)
    add_table(doc, ["기능", "설명", "평가 관점"], [
        ("정책 선택", "사용자가 콤보박스에서 FIFO, Optimal, LRU, Second Chance, LRFU-Lite 선택", "필수 입력 요소 충족"),
        ("참조열 입력", "한 글자를 하나의 page로 해석하는 reference string 입력", "다양한 workload 실험 가능"),
        ("Frame size 입력", "양의 정수 frame 수 입력", "가변 인자 분석 가능"),
        ("단계별 출력", "각 reference마다 Hit/Fault/Migration, victim, frame snapshot 출력", "동작 과정 분석 가능"),
        ("성능 지표", "Hit, Fault, Migration, Fault Rate, 실행 시간, 지연 시간 추정", "정량 비교 가능"),
        ("시각화", "프레임 전이 그리드와 pie chart, 보고서 bar chart 제공", "결과 해석 용이"),
    ], left_cols={1, 2})


def add_background(doc):
    heading(doc, "II. 배경 지식 및 관련 기술", 1)
    heading(doc, "1. 배경 지식", 2)
    heading(doc, "가상 메모리의 필요성", 3)
    add_para(doc, "현대 운영체제에서 메모리 관리는 단순히 물리 RAM을 나누어 주는 문제가 아니다. 하나의 컴퓨터에서는 여러 프로세스가 동시에 실행되고, 각 프로세스는 자신만의 독립적인 주소 공간을 가진 것처럼 동작해야 한다. 만약 모든 프로세스가 실제 물리 주소를 직접 사용한다면 한 프로세스의 잘못된 접근이 다른 프로세스의 데이터를 손상시킬 수 있고, 프로그램을 메모리의 어느 위치에 적재해야 하는지도 복잡해진다. 가상 메모리는 이러한 문제를 해결하기 위해 등장한 운영체제의 핵심 기법이다. 프로세스에는 논리적으로 연속된 큰 주소 공간을 제공하고, 실제 물리 메모리와의 대응 관계는 운영체제와 하드웨어가 관리한다.")
    add_para(doc, "가상 메모리의 장점은 크게 세 가지로 볼 수 있다. 첫째, 프로세스 보호가 가능하다. 각 프로세스는 자신의 가상 주소 공간만 접근한다고 생각하므로 다른 프로세스의 메모리를 직접 침범하기 어렵다. 둘째, 실제 물리 메모리보다 큰 프로그램도 실행할 수 있다. 프로그램 전체를 한 번에 메모리에 올리지 않고 필요한 부분만 적재하면 되기 때문이다. 셋째, 메모리 배치가 유연해진다. 프로세스 입장에서는 연속된 주소처럼 보이지만 실제 물리 frame은 불연속적으로 배치될 수 있다.")

    heading(doc, "Paging과 Page Table", 3)
    add_para(doc, "Paging은 가상 메모리를 고정 크기 단위인 page로 나누고, 물리 메모리를 같은 크기의 frame으로 나누어 관리하는 방식이다. 가상 주소는 page number와 offset으로 나뉜다. Page number는 page table을 통해 물리 frame number로 변환되고, offset은 page 내부 위치를 나타내므로 그대로 사용된다. 예를 들어 page size가 4KB라면 하위 12비트는 offset이고 나머지 상위 비트는 page number가 된다.")
    add_para(doc, "Page table은 가상 page와 물리 frame 사이의 매핑 정보를 저장한다. Page table entry에는 frame number뿐 아니라 valid bit, reference bit, modified bit, protection bit 등이 포함될 수 있다. Valid bit는 해당 page가 현재 물리 메모리에 존재하는지 또는 접근 가능한지를 나타낸다. Reference bit는 최근에 참조되었는지를 나타내며 Second Chance나 Clock 계열 알고리즘에서 사용된다. Modified bit, 또는 dirty bit는 page가 메모리에 올라온 뒤 수정되었는지를 나타낸다. Dirty page를 교체하려면 디스크에 다시 기록해야 하므로 교체 비용이 더 크다.")
    add_para(doc, "Page table 자체도 매우 커질 수 있다. 32비트 또는 64비트 주소 공간에서는 모든 page에 대해 entry를 두면 page table 크기가 커지기 때문에 multi-level page table, inverted page table, segmentation과 paging의 결합 같은 다양한 구조가 사용된다. 강의자료에서도 page table 크기 문제와 segmented paging 개념이 소개된다. 본 프로젝트는 page table 구조 자체를 구현하지는 않지만, page replacement가 page table entry의 valid/reference/modified 정보와 밀접하게 연결된다는 점을 이해하는 것이 중요하다.")

    heading(doc, "Demand Paging", 3)
    add_para(doc, "Demand paging은 page가 실제로 필요해지는 순간에만 메모리에 적재하는 방식이다. 프로그램 실행 시 전체 주소 공간을 모두 물리 메모리에 올리지 않고, CPU가 특정 page를 참조했을 때 그 page가 메모리에 없으면 page fault를 발생시켜 보조기억장치에서 page를 읽어 온다. 이 방식은 초기 적재 시간을 줄이고 물리 메모리를 절약할 수 있다. 사용되지 않는 코드나 데이터는 굳이 메모리에 올리지 않아도 되기 때문이다.")
    add_para(doc, "하지만 demand paging은 page fault 비용이라는 단점을 가진다. Page fault는 일반적인 메모리 접근보다 훨씬 느리다. CPU의 메모리 접근은 매우 빠르지만, page fault가 발생하면 운영체제 trap 처리, 디스크 또는 swap 영역 접근, page table 갱신, 프로세스 재시작 과정이 필요하다. 따라서 demand paging 시스템의 성능은 page fault를 얼마나 줄이느냐에 크게 의존한다. 본 프로젝트의 성능 평가에서 page fault count와 fault rate를 핵심 지표로 둔 이유도 여기에 있다.")

    heading(doc, "Page Fault 처리 과정", 3)
    add_para(doc, "Page fault는 단순히 page가 없다는 사실만을 의미하지 않는다. 운영체제는 먼저 해당 접근이 유효한 주소 접근인지 검사한다. 만약 프로세스가 자신의 주소 공간에 없는 page를 참조했다면 이는 잘못된 메모리 접근이므로 프로세스를 종료하거나 예외를 발생시킨다. 반대로 유효한 page이지만 현재 물리 메모리에 없는 경우라면 운영체제는 해당 page가 보조기억장치의 어디에 있는지 찾고, 이를 적재할 frame을 확보해야 한다.")
    add_para(doc, "빈 frame이 있으면 page fault 처리는 비교적 간단하다. 운영체제는 보조기억장치에서 page를 읽어 빈 frame에 적재하고, page table entry를 valid 상태로 바꾸며 frame number를 기록한다. 그러나 빈 frame이 없으면 기존 frame 중 하나를 선택해 비워야 한다. 이때 선택되는 page가 victim page이고, victim을 고르는 규칙이 page replacement policy이다. Victim page가 dirty page라면 먼저 디스크에 기록해야 하므로 page fault 처리 시간이 더 길어진다.")
    add_para(doc, "Page replacement가 끝나고 새 page가 메모리에 적재되면 운영체제는 page table과 frame table을 갱신하고, page fault를 일으킨 명령을 다시 실행한다. 이 과정에서 잘못된 victim을 선택하면 곧바로 다시 page fault가 발생할 수 있다. 따라서 page replacement policy의 목표는 단순히 빈 frame을 만드는 것이 아니라, 가까운 미래에 다시 필요할 가능성이 낮은 page를 선택하는 것이다.")

    heading(doc, "지역성 원리와 Working Set", 3)
    add_para(doc, "프로그램은 일반적으로 완전히 무작위로 메모리를 참조하지 않는다. 특정 시점에 집중적으로 사용하는 코드와 데이터가 있고, 시간이 지나면 다른 코드와 데이터로 관심 영역이 이동한다. 이 현상을 locality of reference라고 한다. Temporal locality는 최근에 사용된 page가 가까운 미래에 다시 사용될 가능성이 높다는 성질이다. 반복문 내부의 변수, 함수 호출 직후의 stack frame, 최근 접근한 배열 원소 등이 temporal locality의 예가 된다.")
    add_para(doc, "Spatial locality는 어떤 주소가 참조되면 그 주변 주소도 곧 참조될 가능성이 높다는 성질이다. 배열을 순차적으로 탐색하거나 연속된 명령어를 실행하는 경우가 여기에 해당한다. Paging은 page 단위로 데이터를 가져오기 때문에 spatial locality가 존재하면 하나의 page fault로 여러 인접 참조를 처리할 수 있다. 반대로 reference string이 A, B, C, D처럼 서로 다른 page를 한 번씩만 참조한다면 page replacement policy가 좋아도 hit를 만들기 어렵다.")
    add_para(doc, "Working set은 특정 시간 구간 동안 프로세스가 실제로 사용하는 page들의 집합이다. Frame 수가 working set 크기보다 작으면 프로세스는 계속 필요한 page를 쫓아다니며 교체하게 되고, page fault가 급격히 증가한다. 이를 thrashing이라고 한다. 본 프로젝트의 Locality workload는 ABC 작업 집합과 DEF 작업 집합이 전환되는 패턴을 통해 working set 변화가 page replacement에 어떤 영향을 주는지 관찰하기 위해 구성하였다.")

    heading(doc, "Page Replacement 평가 기준", 3)
    add_para(doc, "Page replacement policy의 가장 직접적인 평가 기준은 page fault count이다. 같은 reference string과 frame size에서 page fault가 적을수록 보조기억장치 접근과 운영체제 개입이 줄어든다. 하지만 입력 길이가 다르면 단순 fault count만으로 비교하기 어렵기 때문에 fault rate도 함께 사용한다. Fault rate는 전체 참조 횟수 중 page fault가 발생한 비율이다.")
    add_para(doc, "Migration count도 중요한 보조 지표이다. Page fault 중에서도 빈 frame이 있을 때 발생한 fault와 기존 page를 교체해야 하는 fault는 비용이 다르다. 빈 frame이 있으면 단순 적재만 하면 되지만, migration은 victim page를 제거하고 새 page를 넣는 과정이 포함된다. 실제 운영체제에서는 victim이 dirty page인지 여부에 따라 비용 차이가 더 커진다. 본 프로젝트는 dirty bit를 구현하지 않았지만 migration count를 별도로 출력하여 실제 교체가 얼마나 자주 발생했는지 확인할 수 있게 했다.")
    add_para(doc, "실행 시간은 구현 언어, UI 렌더링, 입력 길이에 영향을 받으므로 작은 실험에서는 page fault count보다 덜 안정적인 지표일 수 있다. 따라서 본 프로젝트에서는 Stopwatch로 실제 시뮬레이션 시간을 측정하되, page fault 비용을 이해하기 위한 보조 지표로 Estimated Delay도 사용하였다. Estimated Delay는 page fault 1회가 10ms의 지연을 만든다고 가정한 단순 모델이다.")

    heading(doc, "Belady's Anomaly와 Stack Property", 3)
    add_para(doc, "일반적인 직관으로는 frame 수가 증가하면 page fault가 감소해야 한다. 더 많은 page를 보관할 수 있으므로 필요한 page가 메모리에 남아 있을 가능성이 높아지기 때문이다. 그러나 FIFO에서는 frame 수가 늘었는데 page fault가 오히려 증가하는 경우가 존재한다. 이 현상을 Belady's anomaly라고 한다. 강의자료의 대표 예제인 1,2,3,4,1,2,5,1,2,3,4,5에서 FIFO는 frame=3일 때 fault 9회, frame=4일 때 fault 10회를 보인다.")
    add_para(doc, "Belady's anomaly가 발생하는 이유는 FIFO가 page의 사용 여부를 고려하지 않고 단순히 적재 시점만 보기 때문이다. Frame이 늘어나면서 큐의 구성과 교체 순서가 달라지고, 그 결과 작은 frame에서는 남아 있던 page가 큰 frame에서는 먼저 제거되는 역설적인 상황이 생긴다. 반면 LRU는 stack property를 만족하는 정책으로 알려져 있다. Frame 수가 증가하면 작은 frame에서 보관하던 page 집합이 큰 frame에서도 포함되는 성질이 있어 Belady's anomaly가 발생하지 않는다.")

    heading(doc, "2. 관련 기술 및 기존 정책", 2)
    heading(doc, "실제 시스템에서의 교체 정책 문제", 3)
    add_para(doc, "관련 기술 파트에서는 본 프로젝트 코드가 각 정책을 어떻게 구현했는지를 설명하기보다, 실제 운영체제와 캐시 시스템에서 왜 page replacement 문제가 어렵고 어떤 기술적 절충이 필요한지를 설명한다. Page replacement는 교재 예제에서는 reference string과 frame size만으로 표현되지만, 실제 시스템에서는 훨씬 많은 정보가 함께 고려된다. 운영체제는 page fault가 발생했을 때 단지 fault 횟수를 줄이는 것뿐 아니라, 디스크 쓰기 비용, 여러 프로세스 간 공정성, working set 유지, 메모리 압박 상황, I/O 중인 page 보호 등을 함께 고려해야 한다.")
    add_para(doc, "실제 운영체제의 page replacement는 전역적인 최적화 문제이기도 하고 지역적인 응답성 문제이기도 하다. 어떤 프로세스가 page fault를 일으켰을 때 그 프로세스의 frame 안에서만 victim을 고르는 local replacement를 사용할 수도 있고, 전체 시스템 frame 중에서 victim을 고르는 global replacement를 사용할 수도 있다. Local replacement는 프로세스 간 간섭을 줄이는 장점이 있지만, 어떤 프로세스는 frame이 부족하고 다른 프로세스는 여유가 있는 상황을 유연하게 처리하기 어렵다. Global replacement는 전체 메모리를 효율적으로 사용할 수 있지만, 한 프로세스가 다른 프로세스의 frame을 빼앗아 성능을 흔들 수 있다.")

    heading(doc, "하드웨어 지원 비트와 운영체제 정책", 3)
    add_para(doc, "운영체제가 page replacement를 수행하려면 page가 최근에 사용되었는지, 수정되었는지, 접근 가능한지 같은 정보를 알아야 한다. 이때 page table entry에 포함될 수 있는 reference bit와 modified bit가 중요하다. Reference bit는 page가 일정 기간 안에 참조되었는지를 나타낸다. Modified bit는 page가 메모리에 올라온 뒤 변경되었는지를 나타낸다. Modified bit가 1인 page를 교체하면 디스크에 다시 기록해야 하므로 비용이 크다. 반면 modified bit가 0인 clean page는 단순히 버리고 나중에 다시 읽어 오면 된다.")
    add_para(doc, "이러한 하드웨어 지원 정보 때문에 실제 정책은 단순히 최근성만 보지 않는다. 예를 들어 같은 정도로 오래 사용되지 않은 page가 두 개 있을 때 하나는 dirty page이고 다른 하나는 clean page라면, 운영체제는 clean page를 먼저 제거하는 편이 전체 I/O 비용을 줄일 수 있다. NUR(Not Used Recently) 또는 Enhanced Second Chance 계열 정책은 reference bit와 modified bit를 함께 사용하여 page를 여러 class로 나누고, 비용이 낮은 class를 우선 교체하려고 한다. 본 프로젝트는 modified bit까지 구현하지는 않았지만, 관련 기술을 이해하면 왜 page fault count 하나만으로 실제 시스템 전체 성능을 완전히 설명하기 어려운지 알 수 있다.")

    heading(doc, "정확한 LRU 구현의 한계와 근사 정책의 필요성", 3)
    add_para(doc, "이론적으로 LRU는 temporal locality를 잘 반영하는 정책이지만, 실제 시스템에서 정확한 LRU를 구현하는 것은 쉽지 않다. 모든 메모리 참조마다 page의 최근성 순서를 정확하게 갱신하려면 하드웨어와 운영체제 모두에 큰 오버헤드가 발생한다. 메모리 참조는 매우 빈번하게 일어나기 때문에, 참조가 발생할 때마다 운영체제가 자료구조를 수정한다면 page fault를 줄이려다가 오히려 일반 실행 비용이 커질 수 있다.")
    add_para(doc, "이 때문에 운영체제에서는 정확한 LRU보다 reference bit를 사용하는 근사 정책이 많이 사용된다. Clock 또는 Second Chance 계열 정책은 page가 최근에 참조되었는지를 한 bit로만 표현한다. 정확한 순서는 잃지만, 최근에 사용된 page를 무작정 제거하지 않는다는 점에서 LRU의 핵심 직관을 낮은 비용으로 흉내 낼 수 있다. 즉 관련 기술의 핵심은 어떤 알고리즘이 이론적으로 가장 좋은가가 아니라, 성능 개선에 필요한 정보를 어느 정도 비용으로 수집하고 유지할 수 있는가이다.")

    heading(doc, "캐시 교체 정책과의 관계", 3)
    add_para(doc, "Page replacement는 운영체제 가상 메모리에서 다루지만, 더 넓게 보면 cache replacement 문제와 같은 계열의 문제이다. CPU cache, 웹 캐시, 데이터베이스 버퍼 캐시, 파일 시스템 page cache는 모두 제한된 공간에 어떤 항목을 남기고 어떤 항목을 제거할지 결정해야 한다. 저장 대상이 page인지 block인지 object인지 다를 뿐, 핵심 질문은 같다. 가까운 미래에 다시 사용할 가능성이 낮은 항목을 어떻게 추정할 것인가이다.")
    add_para(doc, "캐시 교체 분야에서는 최근성(recency), 빈도(frequency), 비용(cost), 크기(size), 만료 시간(expiration) 같은 여러 요소를 결합한다. 예를 들어 웹 캐시에서는 자주 요청되는 object뿐 아니라 object 크기와 가져오는 비용도 중요하다. 데이터베이스 버퍼 캐시에서는 순차 스캔이 캐시를 오염시키지 않도록 별도의 처리 전략을 두기도 한다. 이런 관점에서 보면 page replacement 역시 단순히 FIFO, LRU 중 하나를 고르는 문제가 아니라, workload 특성과 메타데이터 비용을 고려해 적절한 균형점을 찾는 문제이다.")

    heading(doc, "빈도 기반 정책과 오래된 정보 문제", 3)
    add_para(doc, "빈도 기반 정책은 장기적으로 자주 참조되는 page나 object를 보호할 수 있다는 장점이 있다. 그러나 빈도 정보는 시간이 지나도 계속 남기 쉽다. 초반에 많이 참조되었지만 이후에는 사용되지 않는 항목이 높은 count 때문에 계속 남아 있으면, 현재 working set에 필요한 새 항목이 밀려날 수 있다. 이를 cache pollution 또는 stale frequency 문제로 볼 수 있다.")
    add_para(doc, "이 문제를 해결하기 위해 관련 기술에서는 aging, decay, window, admission policy 같은 개념을 사용한다. Aging은 오래된 참조 기록의 영향력을 시간이 지남에 따라 줄이는 방식이다. Window 방식은 최근 일정 구간 안의 참조만 중요하게 본다. Admission policy는 어떤 항목을 캐시에 넣을 것인지부터 신중하게 결정한다. 본 프로젝트의 신규 정책이 감쇠 점수를 사용하는 이유도 이 관련 기술 흐름과 연결된다. 다만 이 장에서는 해당 아이디어의 배경을 설명하고, 실제 LRFU-Lite 구현 절차는 III장에서 별도로 다룬다.")

    heading(doc, "관련 연구와 LRFU 계열의 의미", 3)
    add_para(doc, "LRFU(Least Recently/Frequently Used) 계열 연구는 LRU와 LFU를 완전히 별개의 정책으로 보지 않고, 최근성과 빈도를 하나의 연속적인 spectrum에서 조절할 수 있다고 본다. LRU는 가장 최근 참조에 민감하게 반응하고, LFU는 장기 참조 횟수에 민감하게 반응한다. LRFU 계열의 핵심은 참조 시점과 참조 횟수를 함께 반영하여 page 또는 cache 항목의 가치를 계산하는 것이다.")
    add_para(doc, "이러한 연구는 본 프로젝트의 신규 정책을 설계하는 데 방향을 제공한다. 과제 수준에서 논문의 수식을 완전히 재현하는 것보다 중요한 것은, 수업에서 배운 LRU와 LFU의 장단점을 이해하고 두 정보를 결합하는 합리적인 단순 모델을 만드는 것이다. 따라서 본 보고서에서는 관련 연구를 배경으로 삼고, III장에서는 이를 학부 수준 구현으로 단순화한 LRFU-Lite의 자료구조와 동작 절차를 설명한다.")

    heading(doc, "본 프로젝트와 관련 기술의 연결", 3)
    add_para(doc, "정리하면 관련 기술 파트의 핵심은 실제 시스템에서 page replacement가 단순한 victim 선택 규칙 하나로 끝나지 않는다는 점이다. 하드웨어 bit, dirty page 비용, working set 변화, global/local replacement, 캐시 오염, 오래된 빈도 정보, 근사 정책의 오버헤드가 모두 함께 고려된다. 본 프로젝트는 이러한 복잡한 현실을 모두 구현하지는 않는다. 대신 reference string과 frame size라는 단순한 모델 안에서 주요 정책의 차이를 관찰하고, 관련 연구의 recency-frequency 결합 아이디어를 LRFU-Lite라는 구현 가능한 정책으로 축소한다.")
    add_para(doc, "따라서 II장의 관련 기술 설명은 실제 운영체제와 캐시 시스템의 문제 배경을 제공하고, III장의 본론은 그 배경을 바탕으로 본 프로젝트가 어떤 입력, 자료구조, 알고리즘, 출력 기능으로 구현되었는지를 구체화한다. 이렇게 구분하면 II장은 이론과 선행기술의 맥락, III장은 구현 설계와 기능 명세라는 역할을 갖게 된다.")


def add_main(doc):
    heading(doc, "III. 본론: 페이지 교체 정책 시뮬레이터 설계", 1)
    heading(doc, "1. 주요 주제의 개요", 2)
    add_para(doc, "본 프로젝트의 주요 주제는 가상 메모리 시스템에서 page fault가 발생했을 때 어떤 page를 교체할 것인지 결정하는 page replacement policy를 직접 구현하고 분석하는 것이다. 강의자료에서는 FIFO, Optimal, LRU, Second Chance와 같은 정책을 개념적으로 설명하지만, 실제로 reference string을 하나씩 처리하면서 frame 내부 상태가 어떻게 변하는지 직접 관찰하지 않으면 정책 간 차이를 직관적으로 이해하기 어렵다. 따라서 본 프로젝트는 정책을 단순히 계산식으로 비교하는 것이 아니라, 입력된 reference string의 각 문자마다 현재 frame 상태, hit/fault 여부, victim page, page fault rate를 기록하는 시뮬레이터 형태로 구현하였다.")
    add_para(doc, "시뮬레이터의 입력은 세 가지이다. 첫 번째 입력은 page replacement policy이다. 사용자는 FIFO, Optimal, LRU, Second Chance, LRFU-Lite 중 하나를 선택한다. 두 번째 입력은 reference string이다. 본 프로젝트에서는 과제 요구사항에 맞추어 한 글자를 하나의 page로 해석한다. 예를 들어 123412512345는 1, 2, 3, 4, 1, 2, 5, 1, 2, 3, 4, 5 순서의 page 참조를 의미한다. 세 번째 입력은 frame size이다. Frame size는 동시에 물리 메모리에 유지할 수 있는 page 개수를 의미하며, frame size가 작을수록 교체가 더 자주 발생한다.")
    add_para(doc, "출력은 단순한 최종 fault 수에 그치지 않도록 설계하였다. 각 reference step마다 해당 page가 이미 frame에 있으면 Hit로 기록하고, frame에 없으면 Page Fault로 기록한다. Page Fault가 발생했지만 아직 빈 frame이 있다면 새 page를 빈 frame에 삽입한다. 반대로 frame이 모두 차 있는 상태에서 Page Fault가 발생하면 정책별 victim selection 알고리즘을 실행하여 기존 page 하나를 제거하고 새 page를 삽입한다. 이 경우 상태를 Migration으로 기록한다. 결과적으로 콘솔에는 step 번호, 참조 page, 상태, victim page, frame snapshot이 출력된다.")
    add_para(doc, "본 프로젝트에서 frame snapshot은 중요한 분석 도구이다. 단순히 fault 수만 보면 어떤 정책이 더 좋은지는 알 수 있지만, 왜 그런 결과가 나왔는지는 알기 어렵다. Frame snapshot은 특정 시점에 각 frame에 어떤 page가 들어 있는지를 보여준다. 예를 들어 LRU에서 page 5가 들어올 때 page 3이 교체되었다면, 그 이유는 page 3이 frame에 있는 page 중 가장 오래 전에 참조되었기 때문이다. 이러한 과정을 표로 추적하면 정책의 동작 원리를 단계적으로 설명할 수 있다.")
    add_para(doc, "프로그램은 세 계층으로 나누어 설계하였다. UI 계층은 사용자의 입력을 받고 결과를 화면에 출력한다. 시뮬레이션 계층은 실제 page replacement 정책을 수행한다. 기록 계층은 각 step의 결과를 저장하여 화면 출력과 보고서 분석에 사용할 수 있게 한다. 이 구조는 정책 추가와 결과 출력이 서로 지나치게 섞이지 않도록 하기 위한 것이다. 예를 들어 LRFU-Lite와 같은 새 정책을 추가할 때 UI 전체를 새로 작성할 필요 없이 Core 클래스의 enum과 victim selection 함수만 확장하면 된다.")
    add_para(doc, "또한 본 프로젝트는 과제의 가점 요소인 신규 정책 제안을 포함하기 위해 LRFU-Lite를 구현하였다. LRFU-Lite는 LRU의 최근성 정보와 LFU의 빈도 정보를 하나의 점수로 결합하려는 정책이다. 기존 page들의 점수는 시간이 지날수록 감소하고, 참조된 page는 점수가 증가한다. 따라서 최근에 자주 참조되는 page는 높은 점수를 유지하고, 오래 전에 잠깐 많이 참조되었지만 최근에는 사용되지 않는 page는 점수가 점차 낮아진다. 이 방식은 학부 수준에서 구현할 수 있으면서도 기존 FIFO/LRU와 다른 분석 관점을 제공한다.")
    add_table(doc, ["계층", "구성 파일", "역할"], [
        ("UI 계층", "Form1.cs, Form1.Designer.cs", "사용자 입력 수집, 정책 실행 호출, 콘솔/차트/그리드 출력"),
        ("시뮬레이션 계층", "Core.cs", "정책별 victim 선택, hit/fault/migration 집계, frame snapshot 생성"),
        ("기록 계층", "Page.cs", "각 step의 data, status, victim, location, frameSnapshot 저장"),
        ("실행 진입점", "Program.cs", "WinForms 애플리케이션 시작"),
    ], left_cols={2})
    heading(doc, "입력-처리-출력 흐름", 3)
    add_number(doc, "사용자가 Policy 콤보박스에서 정책을 선택한다.")
    add_number(doc, "사용자가 Reference String에 page 참조열을 입력한다.")
    add_number(doc, "사용자가 #Frame에 물리 frame 수를 입력한다.")
    add_number(doc, "Run 버튼을 누르면 Form1에서 입력값을 검증하고 Core 객체를 생성한다.")
    add_number(doc, "Core는 reference string의 각 문자를 순서대로 Operate 함수에 전달하여 처리한다.")
    add_number(doc, "Operate 함수는 hit 여부를 먼저 검사하고, fault 발생 시 빈 frame 또는 victim 교체를 수행한다.")
    add_number(doc, "각 step 결과는 Page 구조체에 저장되고 pageHistory 목록에 누적된다.")
    add_number(doc, "실행이 끝나면 Form1이 pageHistory를 기반으로 콘솔, frame grid, pie chart, fault rate label을 갱신한다.")
    add_code(doc, "User Input -> Form1 validation -> Core(policy, frameSize)\n"
                  "          -> for each reference: Core.Operate(page, index, referenceString)\n"
                  "          -> Page history + frame snapshot + statistics\n"
                  "          -> Console output + transition grid + chart + report analysis")

    heading(doc, "2. 주요 주제의 핵심 알고리즘 및 기능", 2)
    add_para(doc, "핵심 알고리즘은 모든 정책에서 공통으로 수행되는 기본 처리 절차와, 정책마다 달라지는 victim selection 절차로 나눌 수 있다. 공통 절차는 현재 참조 page가 frame 안에 존재하는지 검사하는 것으로 시작된다. 만약 존재한다면 page fault가 아니므로 hit count만 증가시키고, 정책별 메타데이터를 갱신한다. LRU는 lastUsed 값을 현재 index로 바꾸고, Second Chance는 reference bit를 1로 바꾸며, LRFU-Lite는 모든 score를 감쇠시킨 뒤 참조된 page의 score를 증가시킨다.")
    add_para(doc, "Page가 frame에 존재하지 않는 경우에는 page fault가 발생한다. 이때 frame에 빈 공간이 남아 있다면 victim을 선택할 필요가 없다. 새 page를 빈 frame에 삽입하고 Page Fault 상태로 기록한다. 그러나 frame이 가득 찬 상태라면 migration이 발생한다. Migration은 실제 교체가 일어났다는 뜻이며, 이때 어떤 page를 victim으로 선택하느냐가 정책의 핵심 차이를 만든다.")
    add_para(doc, "본 프로젝트에서 구현한 정책들은 모두 같은 Core.Operate 흐름을 공유한다. 정책별 차이는 SelectVictim 함수 안에서 분기된다. 이렇게 공통 흐름과 정책별 victim 선택을 분리하면 코드 중복이 줄어들고, 정책별 성능 차이도 같은 조건에서 비교할 수 있다. 만약 각 정책을 별도의 완전히 다른 실행 흐름으로 작성하면 통계 계산이나 frame snapshot 생성 방식이 달라져 비교가 불공정해질 수 있다.")
    add_table(doc, ["정책", "Victim 선택 기준", "추가 메타데이터", "시간 복잡도(교체 시)"], [
        ("FIFO", "가장 먼저 들어온 page", "프레임 순서 또는 queue", "O(1) 또는 O(n)"),
        ("Optimal", "미래 참조 시점이 가장 먼 page", "전체 reference string", "O(n * 남은 참조 길이)"),
        ("LRU", "lastUsed 값이 가장 작은 page", "마지막 참조 index", "O(n)"),
        ("Second Chance", "reference bit가 0인 clock hand 위치 page", "reference bit, clock hand", "최악 O(n), 평균적으로 낮음"),
        ("LRFU-Lite", "감쇠 점수가 가장 낮은 page", "score, decay factor", "O(n)"),
    ], left_cols={1, 2})

    heading(doc, "공통 처리 알고리즘", 3)
    add_para(doc, "공통 처리 알고리즘은 reference string의 각 page를 순차적으로 읽는 구조이다. 각 step에서 page가 frame에 있는지 검사하고, 결과에 따라 hit 처리, 빈 frame 삽입, victim 교체 중 하나를 수행한다. 이 구조는 실제 운영체제의 page fault 처리 흐름을 단순화한 것이다. 실제 운영체제에서는 page table lookup과 trap, disk I/O가 포함되지만, 본 프로젝트는 정책 비교가 목적이므로 page 존재 여부와 victim 선택에 집중한다.")
    add_code(doc, "for index from 0 to referenceString.Length - 1:\n"
                  "    page = referenceString[index]\n"
                  "    if page exists in frame:\n"
                  "        status = HIT\n"
                  "        hit++\n"
                  "        update policy metadata\n"
                  "    else:\n"
                  "        fault++\n"
                  "        if frame has empty slot:\n"
                  "            status = PAGEFAULT\n"
                  "            insert page into empty frame\n"
                  "        else:\n"
                  "            status = MIGRATION\n"
                  "            victim = SelectVictim(policy)\n"
                  "            replace victim with page\n"
                  "            migration++\n"
                  "    save frame snapshot and step result")

    heading(doc, "FIFO victim selection", 3)
    add_para(doc, "FIFO는 frame에 먼저 들어온 page를 먼저 제거한다. 샘플 코드도 기본적으로 FIFO 구조를 가지고 있었지만, 정책을 여러 개 추가하기 위해 기존 queue 중심 구조를 List 기반 frame_window와 정책별 메타데이터 구조로 정리하였다. FIFO에서는 frame_window의 가장 앞쪽 page를 victim으로 보고, 해당 page를 제거한 뒤 새 page를 뒤쪽에 추가한다. 이렇게 해야 FIFO의 적재 순서가 유지된다.")
    add_para(doc, "FIFO는 hit가 발생해도 page의 순서를 바꾸지 않는다. 어떤 page가 방금 참조되었더라도 처음 들어온 시점이 오래되었다면 교체 대상이 될 수 있다. 이 점이 LRU와의 핵심 차이이다. 본 프로젝트에서 FIFO가 Textbook workload frame=4에서 fault 10회를 보인 이유도 page의 실제 재참조 가능성을 고려하지 않기 때문이다.")

    heading(doc, "Optimal victim selection", 3)
    add_para(doc, "Optimal은 현재 frame에 들어 있는 각 page가 앞으로 reference string에서 언제 다시 등장하는지 검사한다. 현재 index 이후에 다시 등장하지 않는 page가 있다면 그 page는 가장 좋은 victim이다. 모두 다시 등장한다면 가장 늦게 등장하는 page를 victim으로 선택한다. 이 방식은 미래 정보를 사용하므로 실제 운영체제에서는 직접 사용할 수 없지만, 시뮬레이터에서는 전체 reference string이 입력으로 주어져 있으므로 구현할 수 있다.")
    add_para(doc, "Optimal 구현은 각 replacement 시점마다 frame 안의 모든 page에 대해 reference string의 남은 구간을 검색한다. 따라서 입력 길이가 길고 frame 수가 크면 비용이 커질 수 있다. 그러나 본 프로젝트에서는 알고리즘 비교가 목적이고 reference string 길이가 길지 않으므로, 이해하기 쉬운 선형 탐색 방식으로 구현하였다. Optimal은 다른 정책이 얼마나 이상적인 결과에 가까운지를 판단하는 기준선으로 사용된다.")

    heading(doc, "LRU victim selection", 3)
    add_para(doc, "LRU는 각 page가 마지막으로 참조된 index를 저장한다. Hit가 발생하면 해당 page의 lastUsed 값을 현재 index로 갱신한다. Page fault로 새 page가 들어오면 새 page의 lastUsed도 현재 index로 설정한다. 교체가 필요할 때는 frame 안의 page 중 lastUsed 값이 가장 작은 page를 victim으로 선택한다. 값이 작다는 것은 가장 오래 전에 참조되었다는 의미이다.")
    add_para(doc, "이 구현은 정확한 LRU의 counter 방식과 유사하다. 실제 운영체제에서는 모든 참조마다 정확한 timestamp를 유지하기 어렵기 때문에 LRU approximation을 사용하지만, 시뮬레이터에서는 reference string을 명시적으로 순회하므로 정확한 lastUsed 기록이 가능하다. 따라서 LRU 결과는 Second Chance와 비교할 때 근사 정책이 얼마나 LRU에 가까운지 확인하는 기준으로도 사용된다.")

    heading(doc, "Second Chance victim selection", 3)
    add_para(doc, "Second Chance는 각 page에 reference bit를 둔다. Page가 hit되면 reference bit를 1로 설정한다. 교체가 필요하면 clock hand가 가리키는 page를 확인한다. Reference bit가 0이면 그 page를 victim으로 선택한다. Reference bit가 1이면 해당 bit를 0으로 바꾸고 clock hand를 다음 위치로 이동한다. 이 과정을 victim이 선택될 때까지 반복한다.")
    add_para(doc, "이 방식은 최근 참조된 page를 보호하지만, 정확한 최근 참조 순서를 저장하지는 않는다. 따라서 LRU보다 메타데이터가 적고 구현 비용이 낮다. 본 프로젝트에서는 referenceBits Dictionary와 clockHand 정수 변수를 사용하였다. Replacement가 완료되면 clockHand는 victim 다음 위치로 이동한다. 이 구조는 원형 큐를 도는 Clock 알고리즘과 같은 사고방식이다.")

    heading(doc, "LRFU-Lite victim selection", 3)
    add_para(doc, "LRFU-Lite는 본 프로젝트에서 제안한 신규 정책이다. 이 정책의 핵심은 page의 가치를 하나의 score로 표현하는 것이다. Score는 빈도와 최근성을 동시에 반영한다. 매 reference마다 frame 안에 있는 모든 page의 score를 0.85배로 감쇠시킨다. 그리고 현재 참조된 page가 hit되면 해당 page의 score에 1.0을 더한다. 새로 삽입되는 page의 score는 1.0으로 시작한다.")
    add_para(doc, "이 방식에서 최근에 참조된 page는 방금 1.0이 더해졌기 때문에 높은 점수를 가진다. 자주 반복되는 page는 여러 번 1.0이 더해지므로 역시 높은 점수를 가진다. 반대로 오랫동안 참조되지 않은 page는 매 step마다 0.85배로 감소하므로 점점 낮은 점수를 갖게 된다. 교체가 필요하면 score가 가장 낮은 page를 victim으로 선택한다.")
    add_para(doc, "LRFU-Lite는 LFU의 오래된 빈도 누적 문제를 완화하기 위해 감쇠를 사용한다. 단순 LFU라면 초반에 많이 참조된 page가 이후에도 높은 count를 유지하지만, LRFU-Lite에서는 시간이 지나면 그 영향이 줄어든다. 동시에 LRU와 달리 단 한 번 최근에 참조된 page뿐 아니라 반복적으로 참조된 page도 높은 점수를 유지할 수 있다. 이러한 특성 때문에 Frequency Bias workload에서 A처럼 자주 등장하는 page를 잘 보호할 수 있다.")
    add_code(doc, "LRFU-Lite score update:\n"
                  "    for each resident page p:\n"
                  "        score[p] = score[p] * 0.85\n"
                  "    if current page is resident:\n"
                  "        score[current] = score[current] + 1.0\n"
                  "    if page fault and memory is full:\n"
                  "        victim = page with minimum score")

    heading(doc, "구현 기능 상세", 3)
    add_para(doc, "UI 기능은 과제 요구사항을 직접 반영하도록 구성하였다. Policy 콤보박스에는 구현된 모든 정책이 표시된다. Reference String 입력란은 사용자가 실험하고 싶은 참조열을 직접 넣을 수 있게 한다. #Frame 입력란은 숫자만 입력하도록 KeyPress 이벤트에서 제한하고, 실행 시에도 0 이하 값은 오류로 처리한다. Random 버튼은 임의 reference string을 생성하여 빠른 실험을 도와준다.")
    add_para(doc, "출력 기능은 세 부분으로 나뉜다. 첫째, 텍스트 콘솔은 각 step의 상세 정보를 제공한다. 예를 들어 [7] DATA 5 is Migrated, victim=3, frames=[1 2 5 4]와 같은 형식으로 현재 참조, 상태, victim, frame snapshot을 보여준다. 둘째, 왼쪽 그리드 이미지는 reference string의 진행에 따른 frame 변화를 시각적으로 나타낸다. Page Fault는 빨간색, Migration은 보라색, Hit는 기본 강조 방식으로 구분된다. 셋째, chart는 Hit와 Fault 비율을 pie chart로 표시한다.")
    add_para(doc, "통계 기능은 hit, fault, migration을 별도로 관리한다. Total reference 수는 hit + fault로 계산하고, fault rate는 fault / total x 100으로 계산한다. 실행 시간은 Stopwatch를 이용하여 측정한다. Estimated Delay는 page fault 1회당 10ms 지연이 있다고 가정하여 fault count에 10을 곱한 값이다. 실제 운영체제의 지연 시간은 디스크 종류, swap 상태, dirty page 여부에 따라 달라지지만, 본 프로젝트에서는 정책 간 상대 비교를 위한 보조 지표로 사용하였다.")
    add_para(doc, "저장 기능은 Save 버튼을 통해 현재 frame transition 이미지를 result.jpg로 저장한다. 이를 통해 보고서 작성 시 특정 실행 결과의 시각 자료를 별도로 확보할 수 있다. 본 보고서에서는 프로그램 실행 결과뿐 아니라 별도의 분석 스크립트를 사용하여 전체 workload 결과를 표와 그래프로 정리하였다.")

    heading(doc, "3. 알고리즘별 pseudo code", 2)
    pseudocode = {
        "FIFO": "if page hit: count hit\nelse if free frame exists: insert page\nelse: victim = oldest page in queue; replace victim with new page",
        "Optimal": "if page fault and memory full:\n    for each resident page:\n        find next use after current index\n    victim = page with farthest next use or no future use",
        "LRU": "on every reference: lastUsed[page] = current index\nif replacement needed:\n    victim = page with minimum lastUsed value",
        "Second Chance": "while replacement not done:\n    if referenceBit[clockHand] == 0: replace that page\n    else: referenceBit[clockHand] = 0; advance clockHand",
        "LRFU-Lite": "on every reference:\n    score[p] = score[p] * 0.85 for resident pages\n    if hit: score[page] += 1\n    if fault and full: victim = page with minimum score",
    }
    for name, code in pseudocode.items():
        heading(doc, name, 3)
        add_code(doc, code)
    heading(doc, "4. 수식 및 평가 지표", 2)
    for item in [
        "Hit Count = reference string 처리 중 이미 frame에 존재한 page 참조 횟수",
        "Page Fault Count = 참조한 page가 frame에 없어 적재가 필요한 횟수",
        "Migration Count = page fault 중 빈 frame이 없어 기존 page를 교체한 횟수",
        "Page Fault Rate(%) = Page Fault Count / 전체 참조 횟수 x 100",
        f"Estimated Delay(ms) = Page Fault Count x {DELAY_UNIT_MS}ms, 본 프로젝트에서 비교용으로 둔 단순 지연 모델",
        f"LRFU-Lite Score: S_p(t) = {LRFU_DECAY} x S_p(t-1) + 1(page p가 현재 참조된 경우), 교체 시 min(S_p) 선택",
    ]:
        add_bullet(doc, item)
    heading(doc, "5. 알고리즘 동작 사례", 2)
    ref = "123412512345"
    for policy in POLICIES:
        r = simulate(ref, 4, policy)
        heading(doc, f"{policy} 동작 추적: reference={ref}, frame=4", 3)
        rows = [(x["step"], x["ref"], x["status"], x["victim"], x["snapshot"], x["score"] if policy == "LRFU-Lite" else "-") for x in r["trace"]]
        add_table(doc, ["Step", "Ref", "Status", "Victim", "Frame Snapshot", "Score"], rows, font_size=7.6)
        add_para(doc, f"{policy} 정책의 이 입력에 대한 결과는 hit {r['hit']}회, fault {r['fault']}회, migration {r['migration']}회, fault rate {r['fault_rate']}%이다.")


def add_performance(doc):
    heading(doc, "IV. 성능 평가", 1)
    heading(doc, "1. 실험 환경", 2)
    add_table(doc, ["항목", "내용"], [
        ("구현 언어", "C#"),
        ("프레임워크", ".NET Framework 계열 WinForms"),
        ("개발/실행 파일", "Memory_Policy_Simulator.exe"),
        ("실험 도구", "프로그램 내 시뮬레이터, 검증용 AnalysisData.exe, 보고서 생성용 Python 스크립트"),
        ("제약 사항", "실제 디스크 I/O를 발생시키지 않는 시뮬레이터이므로 지연 시간은 page fault x 10ms의 추정값으로 비교"),
        ("입력 단위", "reference string의 한 글자를 하나의 page로 간주"),
    ], left_cols={1})
    heading(doc, "2. 성능 평가 입력 요소", 2)
    add_table(doc, ["Workload", "Reference String", "선정 이유"], WORKLOADS, left_cols={2})
    add_para(doc, "Frame size는 3, 4, 5로 설정하였다. 이 값들은 강의자료의 예제와 유사하게 작으면서도 page replacement 차이를 관찰하기에 충분하다. 특히 frame size 3과 4는 FIFO에서 Belady's anomaly가 나타나는지 확인하기 좋은 조건이다.")
    heading(doc, "3. 성능 평가 기준", 2)
    for item in [
        "Page Fault Count: 가장 중요한 기준이다. page fault는 디스크 접근과 context switch를 유발할 수 있기 때문이다.",
        "Fault Rate: 입력 길이가 다른 workload를 비교하기 위해 비율로 환산한다.",
        "Migration Count: 빈 frame이 없어 실제 교체가 발생한 횟수이다.",
        "Estimated Delay: page fault의 상대적 비용을 이해하기 위한 보조 지표이다.",
        "Frame Size Sensitivity: frame 수가 증가할 때 정책별 결과가 어떻게 변하는지 확인한다.",
    ]:
        add_bullet(doc, item)
    heading(doc, "4. 실험 결과", 2)
    add_table(doc, ["Workload", "Ref", "Frame", "Policy", "Hit", "Fault", "Migration", "Fault Rate", "Delay(ms)"], [
        (w, ref, f, p, h, fault, m, f"{rate:.2f}%", delay) for w, ref, f, p, h, fault, m, rate, delay in RESULTS
    ], font_size=7.2, left_cols={0, 1, 3})
    heading(doc, "5. 그래프 기반 결과", 2)
    for workload, _, _ in WORKLOADS:
        path = ASSETS / f"{workload.replace(' ', '_').lower()}_frame4.png"
        subset = [(r[3], r[5]) for r in RESULTS if r[0] == workload and r[2] == 4]
        draw_bar_chart(f"{workload} workload, frame=4", subset, path)
        doc.add_picture(str(path), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(doc, "6. Frame size 변화 분석", 2)
    for workload in ["Textbook", "Frequency Bias", "Locality"]:
        path = ASSETS / f"{workload.replace(' ', '_').lower()}_frames.png"
        draw_frame_chart(workload, path)
        doc.add_picture(str(path), width=Inches(6.7))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(doc, "7. 결과 분석", 2)
    analyses = [
        "Textbook workload에서 frame=4 기준 Optimal은 fault 6회로 가장 낮은 결과를 보였다. 이는 미래 참조 정보를 알고 있어 가장 늦게 다시 사용될 page를 제거할 수 있기 때문이다. LRU와 Second Chance는 fault 8회로 FIFO 10회보다 우수했다. FIFO는 page의 실제 재사용 가능성을 고려하지 않으므로 오래 머문 page를 제거하다가 가까운 미래에 다시 필요한 page를 잃을 수 있다.",
        "Textbook workload에서 FIFO는 frame=3일 때 fault 9회, frame=4일 때 fault 10회가 발생했다. frame 수가 늘었는데 fault가 증가했으므로 Belady's anomaly가 확인된다. 이는 FIFO가 stack property를 만족하지 않기 때문에 발생한다. 반면 LRU는 frame 수가 증가할수록 보유 가능한 최근 참조 집합이 확장되므로 일반적으로 fault가 감소하거나 유지된다.",
        "Sequential Scan workload에서는 모든 정책의 fault rate가 100%이다. A부터 L까지 각 page가 한 번만 등장하므로 어떤 page를 보존해도 다음 참조에서 hit가 발생하지 않는다. 이 결과는 교체 정책 자체보다 workload의 재참조 특성이 더 중요할 수 있음을 보여준다.",
        "Locality workload는 A/B/C 집합과 D/E/F 집합이 번갈아 나타난다. LRU는 현재 작업 집합을 빠르게 따라가는 편이지만, LRFU-Lite는 이전에 자주 등장한 page의 점수를 한동안 보존한다. 따라서 감쇠 계수가 너무 크면 작업 집합 전환 시 오래된 빈도 정보가 방해가 될 수 있다.",
        "Frequency Bias workload에서는 A가 매우 자주 반복된다. frame=3에서 FIFO, LRU, Second Chance는 fault 7회인 반면 LRFU-Lite와 Optimal은 fault 6회를 보였다. 이 경우 LRFU-Lite는 반복적으로 등장하는 A의 점수를 높게 유지하여 A를 보호하고, 빈도가 낮은 page를 교체 대상으로 선택하는 효과를 보였다.",
    ]
    for text in analyses:
        add_para(doc, text)
    heading(doc, "8. 신규 정책 LRFU-Lite의 세부 평가", 2)
    add_para(doc, "LRFU-Lite는 LRU와 LFU의 장단점을 절충하기 위해 설계하였다. LRU는 최근성에 민감하지만 장기적으로 자주 쓰이는 page를 충분히 보호하지 못할 수 있다. LFU는 빈도를 잘 반영하지만 오래전에 많이 사용된 page가 계속 남는 문제가 있다. LRFU-Lite는 점수 감쇠를 통해 오래된 빈도 정보를 서서히 약화시키고, 최근 참조가 발생하면 점수를 즉시 증가시킨다.")
    add_para(doc, "이 정책은 학부 수준 구현 난이도에 맞게 단순한 Dictionary<char,double> 자료구조로 구현하였다. 매 참조마다 resident page 수만큼 점수를 갱신하므로 교체 대상 탐색까지 포함해 O(n) 수준이다. n은 frame size이므로 본 프로젝트와 같은 작은 frame 실험에서는 충분히 이해 가능하고 구현도 간단하다.")
    add_para(doc, "단점도 분명하다. decay factor를 고정하면 모든 workload에 최적으로 동작하기 어렵다. 0.85는 빈도 정보를 어느 정도 유지하면서도 시간이 지나면 약화시키기 위한 중간값이다. 그러나 작업 집합 전환이 빠른 workload에서는 더 작은 감쇠 계수가 필요하고, 장기 인기 page가 중요한 workload에서는 더 큰 감쇠 계수가 유리할 수 있다.")


def add_conclusion_refs_appendix(doc):
    heading(doc, "V. 결론", 1)
    add_para(doc, "본 프로젝트는 가상 메모리 페이지 교체 정책의 동작을 직접 구현하고 실험적으로 비교한 결과물이다. 기존 FIFO 샘플 프로그램을 확장하여 Optimal, LRU, Second Chance, LRFU-Lite를 추가하였고, 정책별 단계 상태와 통계 지표를 출력하도록 개선하였다. 이 과정에서 page replacement가 단순한 암기 주제가 아니라 workload 특성과 메타데이터 설계에 따라 성능이 달라지는 문제임을 확인하였다.")
    add_para(doc, "핵심 결론은 네 가지이다. 첫째, Optimal은 이론적 기준선으로 가장 낮은 fault를 보이지만 실제 시스템에서는 미래 참조열을 알 수 없으므로 비교 기준으로 사용하는 것이 적절하다. 둘째, FIFO는 구현이 매우 단순하지만 Belady's anomaly가 나타날 수 있어 일반적인 성능 안정성은 낮다. 셋째, LRU와 Second Chance는 시간 지역성을 반영하여 많은 경우 FIFO보다 나은 결과를 보인다. 넷째, LRFU-Lite는 빈도 편향 workload에서 효과를 보였지만 감쇠 계수 선택이 중요하다.")
    add_para(doc, "향후 개선 방안으로는 LRFU-Lite의 decay factor를 사용자 입력으로 받도록 확장하고, 0.50, 0.70, 0.85, 0.95 등 여러 계수에 대해 workload별 성능을 비교할 수 있다. 또한 dirty bit와 write-back 비용을 추가하면 Enhanced Second Chance나 NUR 계열 정책까지 평가할 수 있다. 마지막으로 여러 프로세스의 local/global replacement를 구분하면 실제 운영체제와 더 가까운 모델로 확장할 수 있다.")
    heading(doc, "참고 자료", 1)
    refs = [
        ("[1]", "운영체제 강의자료 Ch3. Memory Management and Virtual Memory, Demand Paging 및 Page Replacement 부분."),
        ("[2]", "Term Project - Page Replacement Policy Design.pdf, 영남대학교 컴퓨터공학과 운영체제 Term Project 과제 안내."),
        ("[3]", "Page replacement algorithm, Wikipedia, https://en.wikipedia.org/wiki/Page_replacement_algorithm"),
        ("[4]", "Least frequently used, Wikipedia, https://en.wikipedia.org/wiki/Least_frequently_used"),
        ("[5]", "Cache replacement policies, Wikipedia, https://en.wikipedia.org/wiki/Cache_replacement_policies"),
        ("[6]", "D. Lee, J. Choi, J. Kim, S. H. Noh, S. L. Min, Y. Cho, and C. S. Kim, LRFU: A Spectrum of Policies that Subsumes the Least Recently Used and Least Frequently Used Policies, IEEE Transactions on Computers, 2001."),
        ("[7]", "Dhruv Matani, Ketan Shah, Anirban Mitra, An O(1) algorithm for implementing the LFU cache eviction scheme, arXiv, https://arxiv.org/abs/2110.11602"),
        ("[8]", "Gil Einziger, Roy Friedman, Ben Manes, TinyLFU: A Highly Efficient Cache Admission Policy, arXiv, https://arxiv.org/abs/1512.00727"),
    ]
    add_table(doc, ["번호", "자료"], refs, left_cols={1})
    doc.add_page_break()
    heading(doc, "부록 A. 프로그램 사용 설명", 1)
    for item in [
        "Memory_Policy_Simulator.exe를 실행한다.",
        "Policy 콤보박스에서 정책을 선택한다.",
        "Reference String에 참조열을 입력한다. 예: 123412512345 또는 ABCDABEFABGHABCD.",
        "#Frame에 프레임 수를 입력한다. 예: 3, 4, 5.",
        "Run 버튼을 누르면 오른쪽 콘솔에 단계별 상태와 최종 통계가 출력된다.",
        "Save 버튼을 누르면 현재 프레임 전이 이미지가 result.jpg로 저장된다.",
    ]:
        add_number(doc, item)
    heading(doc, "부록 B. 핵심 코드 설명", 1)
    add_para(doc, "Core.cs의 Operate 함수는 하나의 page reference를 처리하는 핵심 함수이다. 먼저 현재 page가 frame_window에 존재하는지 검사한다. 존재하면 Hit로 기록하고 lastUsed, frequency, reference bit, LRFU score를 갱신한다. 존재하지 않으면 Page Fault를 증가시키고, 빈 frame이 있으면 단순 삽입, 빈 frame이 없으면 SelectVictim 함수로 victim 위치를 선택한다.")
    add_code(doc, "public Page.STATUS Operate(char data, int index, string referenceString)\n{\n    if (frame contains data) update hit metadata;\n    else if (free frame exists) insert new page;\n    else select victim according to policy and replace;\n    save frameSnapshot into pageHistory;\n}")
    add_para(doc, "SelectVictim 함수는 policy enum에 따라 FIFO, Optimal, LRU, Second Chance, LRFU-Lite 중 하나의 victim selection 함수를 호출한다. 이 구조를 사용하면 새 정책을 추가할 때 enum 값과 victim selection 함수만 추가하면 되므로 확장성이 좋다.")
    heading(doc, "부록 C. 추가 단계별 추적: Frequency Bias, frame=3", 1)
    ref = "AAAABCAAADEFAAA"
    for policy in POLICIES:
        r = simulate(ref, 3, policy)
        heading(doc, f"{policy}: hit={r['hit']}, fault={r['fault']}, migration={r['migration']}", 2)
        rows = [(x["step"], x["ref"], x["status"], x["victim"], x["snapshot"], x["score"] if policy == "LRFU-Lite" else "-") for x in r["trace"]]
        add_table(doc, ["Step", "Ref", "Status", "Victim", "Frame Snapshot", "Score"], rows, font_size=7.4)
    heading(doc, "부록 D. 전체 실험 결과 원자료", 1)
    add_table(doc, ["Workload", "Ref", "Frame", "Policy", "Hit", "Fault", "Migration", "Fault Rate", "Delay(ms)"], [
        (w, ref, f, p, h, fault, m, f"{rate:.2f}%", delay) for w, ref, f, p, h, fault, m, rate, delay in RESULTS
    ], font_size=7.0, left_cols={0, 1, 3})


def build():
    doc = Document()
    setup_doc(doc)
    add_title_block(doc)
    add_summary(doc)
    doc.add_page_break()
    add_intro(doc)
    add_background(doc)
    add_main(doc)
    add_performance(doc)
    add_conclusion_refs_appendix(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
