import os
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_narrative_report import (  # noqa: E402
    ASSETS,
    FRAMES,
    POLICIES,
    RESULTS,
    WORKLOADS,
    add_graph_pair,
    bullet,
    code,
    metric_table,
    para,
    set_font,
    setup_doc,
    simulate,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "대표학번_가상메모리_페이지교체정책_보고서_양식반영_상세서술.docx"
if os.environ.get("REPORT_OUT"):
    OUT = Path(os.environ["REPORT_OUT"])


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("가상 메모리 페이지 교체 정책 설계")
    set_font(r, size=22, bold=True, color="1F3A5F")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("운영체제 Term Project 보고서")
    set_font(r, size=13, italic=True, color="555555")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("소속: 컴퓨터공학과    학번: 제출 전 입력    이름: 제출 전 입력")
    set_font(r, size=10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("구현 언어 및 환경: C# WinForms / Visual Studio / 제출 Release 실행파일 기준")
    set_font(r, size=9.5, color="555555")


def add_summary(doc):
    doc.add_heading("[요 약]", level=1)
    para(doc, "본 프로젝트에서는 수업 시간에 학습한 가상 메모리의 페이지 교체 정책을 직접 구현하고, 같은 reference string에 대해 정책별 동작 결과가 어떻게 달라지는지 분석하였다. 제공된 FIFO 기반 C# WinForms 시뮬레이터를 확장하여 FIFO, NUR (0,1 first), NUR (1,0 first), Second Chance, LRFU-Lite 정책을 선택할 수 있도록 하였고, 각 참조 단계마다 frame 상태와 victim page, 정책별 내부 상태를 확인할 수 있게 하였다.")
    para(doc, "가상 메모리에서 page fault는 프로그램 실행 중 필요한 page가 물리 메모리에 없을 때 발생한다. 강의자료 Ch3에서 설명하듯이 물리 메모리가 가득 찬 상태에서는 page replacement algorithm을 통해 victim frame을 선택해야 하며, 이 선택이 전체 page fault 수와 성능 지연에 큰 영향을 준다. 본 보고서는 이러한 이론적 배경을 바탕으로, 실제 구현한 정책들이 어떤 기준으로 victim을 선택하는지와 그 결과가 왜 달라지는지를 단계별로 설명한다.")
    para(doc, "실험은 Textbook, NURPriority, ClockSensitive, FrequencyBias, LocalityShift 다섯 reference string과 frame size 3, 4, 5를 사용하였다. 또한 Second Chance의 clock start, NUR의 R reset interval과 Modified Pages 입력을 바꿔 보면서 정책 고유 매개변수가 결과에 미치는 영향도 확인하였다. 분석 결과, NUR 계열은 R/M bit class와 reset 주기에 민감하게 반응했고, Second Chance는 clock hand의 시작 위치에 따라 결과가 달라졌다. 신규 정책으로 제안한 LRFU-Lite는 반복 참조가 강한 입력에서는 효과적이었지만, 작업 집합이 갑자기 바뀌는 입력에서는 과거 score가 남아 오히려 불리해질 수 있었다.")
    para(doc, "키워드: 가상 메모리, Demand Paging, Page Fault, Page Replacement, FIFO, NUR, Second Chance, Clock Algorithm, LRFU-Lite, Reference Bit, Modified Bit")


def add_intro(doc):
    doc.add_heading("I. 서론", level=1)
    doc.add_heading("1. 프로젝트 배경과 개요", level=2)
    para(doc, "운영체제는 여러 프로그램이 동시에 실행되는 것처럼 보이게 하면서도, 각 프로그램이 자기만의 독립적인 주소 공간을 가진다고 느끼게 해야 한다. 이를 위해 사용되는 핵심 기법이 가상 메모리이다. 가상 메모리를 사용하면 프로그램 전체를 처음부터 물리 메모리에 올리지 않아도 되고, 실제로 참조되는 page만 필요할 때 적재할 수 있다. 이 방식은 메모리 사용량을 줄이고 큰 프로그램도 실행할 수 있게 해 주지만, 필요한 page가 메모리에 없을 때 page fault가 발생한다.")
    para(doc, "Page fault가 발생했을 때 빈 frame이 있으면 새 page를 넣으면 된다. 그러나 물리 메모리가 이미 가득 차 있으면 기존 page 중 하나를 골라 내보낸 뒤 새 page를 적재해야 한다. 이때 어떤 page를 victim으로 선택할 것인지가 페이지 교체 정책의 핵심이다. 잘못된 page를 내보내면 곧바로 다시 필요한 page가 없어져 page fault가 반복될 수 있고, 실제 시스템에서는 디스크 접근과 문맥 처리 비용이 커진다.")
    para(doc, "본 프로젝트는 이 과정을 단순한 설명으로 끝내지 않고 직접 시뮬레이터로 구현한 뒤, 여러 입력에 대해 결과를 비교하는 데 목적이 있다. 사용자는 페이지 교체 정책, reference string, frame size, clock start, R reset interval, modified pages를 입력하고, 프로그램은 각 step의 hit, page fault, migration 여부와 현재 frame 상태를 출력한다. 따라서 최종 fault 수뿐 아니라 정책이 왜 그런 결정을 했는지도 확인할 수 있다.")

    doc.add_heading("2. 프로젝트 목표 및 의의", level=2)
    para(doc, "첫 번째 목표는 수업에서 배운 demand paging과 page replacement 개념을 실제 코드로 옮겨 보는 것이다. 강의자료에서는 valid-invalid bit, page fault 처리 과정, free frame이 없을 때 victim frame을 선택하는 과정, dirty page를 고려하는 이유 등을 설명한다. 이 프로젝트에서는 그중 page replacement 부분을 중심으로, frame 배열과 정책별 메타데이터를 직접 관리하였다.")
    para(doc, "두 번째 목표는 FIFO만 구현하는 데서 멈추지 않고, 최근 사용 여부와 수정 여부, clock hand, score를 활용하는 정책들을 비교하는 것이다. FIFO는 구현이 쉽지만 최근 사용 정보를 반영하지 못한다. NUR는 reference bit와 modified bit를 사용하고, Second Chance는 reference bit와 clock hand를 사용하며, LRFU-Lite는 page별 score를 사용한다. 이렇게 서로 다른 정보를 사용하는 정책을 같은 입력에서 비교하면 정책의 장단점이 더 분명하게 보인다.")
    para(doc, "세 번째 목표는 새로운 정책을 제안하고 그 한계까지 함께 분석하는 것이다. LRFU-Lite는 반복적으로 참조되는 page를 보호하기 위해 score를 누적하고, 시간이 지날수록 score를 감쇠시키는 방식이다. 단순히 신규 정책이 좋다고 주장하는 것이 아니라, 어떤 workload에서는 좋은 결과가 나오고 어떤 workload에서는 좋지 않은 결과가 나오는지를 함께 분석하였다.")

    doc.add_heading("3. 구현하고자 하는 핵심 기능", level=2)
    para(doc, "본 프로그램의 핵심 기능은 다섯 가지이다. 첫째, 사용자가 선택한 정책에 따라 같은 reference string을 서로 다른 방식으로 처리한다. 둘째, 각 page 참조마다 frame snapshot을 저장하여 동작 순서를 확인할 수 있게 한다. 셋째, hit count, page fault count, page fault rate, migration count, estimated delay를 출력한다. 넷째, NUR와 Second Chance처럼 추가 입력이 필요한 정책을 위해 reset interval, modified pages, clock start를 제공한다. 다섯째, 결과를 grid와 pie chart로 시각화하여 단순 텍스트보다 쉽게 비교할 수 있게 한다.")


def add_background(doc):
    doc.add_heading("II. 배경 지식 및 관련 기술", level=1)
    doc.add_heading("1. 배경 지식", level=2)
    para(doc, "강의자료 Ch3의 가상 메모리 부분에서는 프로그램 크기와 실제 물리 메모리 크기가 항상 일대일로 맞지 않는다는 점을 설명한다. 예를 들어 프로그램은 큰 주소 공간을 사용하는 것처럼 보이지만, 실제 물리 메모리에는 필요한 일부 page만 올라와 있을 수 있다. 이때 디스크나 swap 영역은 물리 메모리의 확장처럼 사용되고, 필요한 page가 참조되는 순간 메모리로 들어온다.")
    para(doc, "Demand paging에서는 page가 실제로 참조될 때 메모리 적재 여부를 확인한다. Page table entry의 valid-invalid bit가 해당 page가 메모리에 있는지 판단하는 데 사용되며, valid bit가 0이면 page fault가 발생한다. Page fault가 발생하면 운영체제는 필요한 page가 어디에 있는지 찾고, 빈 frame이 있으면 그곳에 적재하고, 빈 frame이 없으면 page replacement algorithm으로 victim frame을 선택한다.")
    para(doc, "강의자료의 Basic Page Replacement 과정은 크게 네 단계로 볼 수 있다. 먼저 필요한 page가 디스크의 어디에 있는지 찾고, 다음으로 free frame이 있는지 확인한다. Free frame이 있으면 바로 사용하지만, 없으면 page replacement algorithm을 이용해 victim frame을 선택한다. 그런 다음 새 page를 frame에 읽어 오고, page table과 frame table을 갱신한 뒤 중단되었던 명령을 다시 실행한다. 본 프로젝트의 시뮬레이터는 이 전체 운영체제 동작을 모두 구현하지는 않지만, 그중 victim frame 선택 과정을 집중적으로 다룬다.")
    para(doc, "Page fault는 가상 메모리 시스템에서 중요한 성능 요인이다. 일반적인 메모리 접근보다 훨씬 큰 비용을 발생시키기 때문이다. 강의자료에서도 page fault를 줄이기 위해 좋은 page replacement algorithm이 필요하다고 설명한다. 본 프로젝트의 estimated delay는 page fault 1회당 10ms라는 단순한 비용을 곱한 값인데, 실제 디스크 시간과 완전히 같지는 않지만 fault 수 증가가 성능 저하로 이어진다는 점을 보여주기 위해 사용하였다.")
    para(doc, "이때 hit와 page fault를 구분하는 것도 중요하다. Hit는 참조한 page가 이미 frame 안에 있어서 추가 적재가 필요 없는 경우이다. Page fault는 참조한 page가 frame 안에 없어서 새로 가져와야 하는 경우이다. 본 프로그램에서는 page fault 중에서도 frame이 비어 있어 단순히 넣기만 한 경우와, 이미 frame이 가득 차서 기존 page를 내보내야 하는 경우를 나누어 본다. 후자를 migration으로 기록하였다.")
    para(doc, "강의자료에서는 modified bit 또는 dirty bit를 이용하면 page transfer overhead를 줄일 수 있다고 설명한다. 수정되지 않은 page는 디스크에 다시 쓸 필요 없이 버릴 수 있지만, 수정된 page는 write-back 비용이 생길 수 있다. 본 프로젝트의 NUR 정책은 이 개념을 단순화하여 Modified Pages 입력에 포함된 page를 M=1로 처리하고, R/M bit 조합에 따라 victim 우선순위를 다르게 둔다.")
    para(doc, "Reference bit는 page가 최근에 참조되었는지를 보여 주는 단순한 정보이다. 정확한 마지막 참조 시각을 모두 저장하는 방식보다 정보량은 적지만, 구현 비용이 낮다. 강의자료에서도 reference bit를 이용한 근사 방법의 장점과 한계를 설명한다. 본 프로젝트의 NUR와 Second Chance는 모두 reference bit를 사용하지만, 사용하는 방식은 다르다. NUR는 전체 resident page의 R/M class를 비교하고, Second Chance는 clock hand가 만나는 순서대로 R bit를 확인한다.")
    para(doc, "또 하나 중요한 개념은 locality와 working set이다. 프로그램은 완전히 무작위로 page를 참조하지 않고, 최근 사용한 page나 가까운 범위의 page를 다시 사용할 가능성이 높다. Working set은 일정 구간 동안 실제로 자주 사용하는 page 집합을 의미하며, frame 수가 working set보다 부족하면 page fault가 계속 증가할 수 있다. 본 실험의 LocalityShift 입력은 작업 집합이 ABC에서 DEF로 바뀌는 상황을 만들기 위해 사용하였다.")
    para(doc, "Thrashing도 이 배경과 연결된다. 강의자료에서는 프로세스가 충분한 page frame을 가지지 못하면 page-fault rate가 높아지고, 실제 계산보다 page를 swap in/out하는 데 많은 시간을 쓰게 된다고 설명한다. 본 프로젝트는 여러 프로세스를 다루지는 않지만, frame size가 작을 때 fault가 늘어나는 현상을 통해 비슷한 문제를 작은 규모로 관찰할 수 있다.")

    doc.add_heading("2. 관련 기술", level=2)
    para(doc, "FIFO는 가장 먼저 들어온 page를 가장 먼저 내보내는 정책이다. 구현이 단순하고 추가 메타데이터가 거의 필요하지 않지만, 강의자료의 FIFO 예시처럼 더 많은 frame을 주었는데도 page fault가 늘어나는 Belady anomaly가 발생할 수 있다. 본 프로젝트의 Textbook 입력도 강의자료의 대표 reference string을 바탕으로 하며, FIFO가 최근 사용 정보를 반영하지 못하는 문제를 보여준다.")
    para(doc, "FIFO의 단점은 '오래 있었다'는 사실과 '앞으로 필요하지 않다'는 사실이 같지 않다는 데 있다. 어떤 page는 오래 전에 들어왔지만 계속 반복해서 사용될 수 있다. FIFO는 hit가 발생해도 그 page의 위치를 바꾸거나 보호 표시를 남기지 않는다. 그래서 실험에서 반복 참조가 있는 입력을 넣으면 FIFO가 불필요한 교체를 만드는 장면을 쉽게 볼 수 있다.")
    para(doc, "Second Chance는 FIFO의 단순함에 reference bit를 결합한 정책이다. Clock order로 page를 검사하면서 R bit가 1이면 해당 bit를 0으로 내리고 page를 그대로 둔다. R bit가 0인 page를 만나면 그 page를 victim으로 선택한다. 강의자료에서는 이를 Clock Replacement라고도 설명하며, circular queue 형태로 구현할 수 있다고 한다. 본 구현도 clockHand를 사용하여 frame을 순환하면서 victim을 찾는다.")
    para(doc, "Second Chance는 이름 그대로 page에게 한 번 더 기회를 주는 방식이다. 하지만 이 기회는 조건부이다. R bit가 1인 page는 한 번 지나갈 수 있지만, 그 과정에서 R bit가 0으로 바뀐다. 이후 다시 참조되지 않으면 다음 scan에서 victim이 될 수 있다. 따라서 Second Chance는 최근 참조 page를 무조건 보호하는 정책이 아니라, FIFO 순서를 유지하면서 reference bit를 이용해 제거 시점을 조금 늦추는 정책으로 이해할 수 있다.")
    para(doc, "NUR는 Not Used Recently의 약자로, reference bit와 modified bit를 함께 사용한다. Page는 (R,M) 조합에 따라 (0,0), (0,1), (1,0), (1,1) class로 나뉜다. 강의자료에서는 (0,1)과 (1,0) 중 어느 쪽을 먼저 제거할지 고민해야 한다고 설명한다. 본 프로젝트는 이 고민을 직접 비교하기 위해 NUR (0,1 first)와 NUR (1,0 first) 두 변형을 구현하였다.")
    para(doc, "NUR에서 (0,0)은 최근에 사용되지 않았고 수정되지도 않은 page이므로 가장 제거하기 쉽다. (1,1)은 최근 사용되었고 수정되었으므로 제거 부담이 가장 크다. 문제는 (0,1)과 (1,0)이다. (0,1)은 최근 사용되지 않았지만 modified page이고, (1,0)은 최근 사용되었지만 clean page이다. 최근성만 보면 (0,1)을 먼저 제거하는 것이 자연스럽고, write-back 비용만 보면 (1,0)을 먼저 제거하는 것도 의미가 있다.")
    para(doc, "LRFU-Lite는 본 프로젝트에서 추가로 제안한 정책이다. 강의자료에서 직접 제공된 표준 알고리즘은 아니지만, 최근 참조와 반복 참조를 모두 반영하고 싶다는 생각에서 설계하였다. 각 resident page는 score를 가지고, 매 step마다 score가 0.85배로 감소한다. 참조된 page는 score가 1.0 증가한다. Replacement가 필요할 때는 score가 가장 낮은 page를 victim으로 선택한다.")
    para(doc, "LRFU-Lite를 제안한 이유는 reference bit만으로는 반복 참조의 강도를 충분히 표현하기 어렵다고 보았기 때문이다. R bit는 최근 참조 여부만 0 또는 1로 나타낸다. 반면 score는 여러 번 참조된 page가 더 높은 값을 가지도록 만들 수 있다. 물론 score가 오래 남으면 작업 집합이 바뀌었을 때 적응이 늦어질 수 있으므로, decay factor를 함께 둔 것이다.")


def add_main(doc):
    doc.add_heading("III. (본론) 가상 메모리 페이지 교체 정책 시뮬레이터 설계", level=1)
    doc.add_heading("1. 주요 주제의 개요", level=2)
    para(doc, "본 프로젝트의 주요 주제는 page replacement policy를 선택할 수 있는 시뮬레이터를 구현하고, 정책별 동작 차이를 관찰하는 것이다. 프로그램은 WinForms 기반 GUI를 사용하며, 사용자는 reference string과 frame size를 입력한 뒤 정책을 선택한다. Reference string은 과제 조건에 맞게 한 글자 단위 page로 처리된다. 예를 들어 ABCDAB은 A, B, C, D, A, B 순서로 page를 참조하는 입력이다.")
    para(doc, "프로그램의 중심은 Core.cs의 Core 클래스이다. Form1.cs는 사용자 입력을 읽고, Core 객체를 만든 뒤 reference string의 각 문자를 순서대로 Core.Operate에 넘긴다. Core.Operate는 현재 page가 frame 안에 있는지 확인하고, hit 또는 page fault를 기록하며, 필요하면 정책별 victim selection을 수행한다. 각 step의 결과는 Page 구조체에 저장되어 pageHistory에 쌓인다.")
    para(doc, "출력은 크게 세 종류로 나뉜다. 첫째, 콘솔 영역에는 step 번호, 참조 page, hit/page fault/migration 여부, victim, frame snapshot, 정책별 내부 상태가 출력된다. 둘째, 그림 영역에는 frame 변화가 grid 형태로 표시된다. 셋째, chart와 label에는 hit/fault 비율과 page fault rate가 표시된다. 이 구조 덕분에 사용자는 숫자 결과와 동작 순서를 함께 볼 수 있다.")
    para(doc, "Core 클래스가 관리하는 주요 상태는 frame_window, pageHistory, referenceBits, modifiedBits, lrfuScores, clockHand이다. frame_window는 현재 메모리에 올라와 있는 page 목록이고, pageHistory는 각 참조 시점의 결과를 저장한다. referenceBits와 modifiedBits는 NUR와 Second Chance에서 사용되고, lrfuScores는 LRFU-Lite에서 사용된다. clockHand는 Second Chance가 다음에 어느 frame부터 검사할지 나타낸다.")
    para(doc, "Page 구조체에는 단순한 page 문자만 저장되는 것이 아니라, status, loc, victim, frameSnapshot, detail, algorithmState가 저장된다. 이 구조 때문에 프로그램은 최종 결과뿐 아니라 각 step의 동작 이유를 설명할 수 있다. 예를 들어 NUR에서는 algorithmState에 R/M bit 상태가 남고, Second Chance에서는 clock 위치와 scan 결과가 남으며, LRFU-Lite에서는 score snapshot이 남는다.")
    para(doc, "실제 구현된 정책은 FIFO, NUR (0,1 first), NUR (1,0 first), Second Chance, LRFU-Lite이다. 보고서에서는 이 다섯 정책만 분석 대상으로 삼았다. 과제 설명에는 여러 정책 후보가 제시되어 있지만, 제출 실행파일과 소스에 들어간 정책을 기준으로 보고서를 작성해야 실제 구현과 문서 내용이 맞는다.")
    para(doc, "UI 측면에서는 정책 선택 콤보박스, reference string 입력창, frame size 입력창, clock start 입력창, R reset interval 입력창, modified pages 입력창을 제공한다. Frame size와 clock start는 양의 정수여야 하고, reset interval은 0 이상의 정수여야 한다. 이러한 입력 검사는 잘못된 실험 조건으로 인해 프로그램이 비정상적으로 동작하는 것을 막아 준다.")

    doc.add_heading("2. 주요 주제의 핵심 알고리즘 및 기능", level=2)
    para(doc, "모든 정책은 공통적으로 같은 실행 흐름을 가진다. 먼저 R bit reset이 필요한지 확인하고, LRFU-Lite라면 resident page의 score를 감쇠시킨다. 그 다음 현재 page가 frame 안에 있는지 찾는다. Page가 있으면 hit로 기록하고, 없으면 page fault로 기록한다. Frame에 빈 자리가 있으면 단순 삽입이고, 빈 자리가 없으면 정책별 victim을 선택한 뒤 migration으로 기록한다.")
    code(doc, "for each page in referenceString:\n    if reset interval condition is met:\n        clear reference bits of resident pages\n    if policy is LRFU-Lite:\n        decay scores of resident pages\n    if page is already in frame:\n        count hit and update metadata\n    else:\n        count page fault\n        if free frame exists:\n            insert page\n        else:\n            select victim by policy\n            replace victim and count migration\n    save frame snapshot and policy state")
    para(doc, "FIFO는 frame_window의 맨 앞 page를 victim으로 선택한다. Hit가 발생해도 frame 순서가 바뀌지 않는다. 그래서 구현은 가장 쉽지만, 최근에 사용된 page라도 오래 전에 들어왔다는 이유로 제거될 수 있다. Core.cs에서는 FIFO일 때 victimIndex를 0으로 두고 RemoveAt(0) 후 새 page를 Add하는 방식으로 queue처럼 처리한다.")
    code(doc, "if policy == FIFO:\n    victimIndex = 0\n    remove the oldest page from frame_window\n    add the new page at the end of frame_window")
    para(doc, "NUR 두 변형은 R bit와 M bit를 보고 victim을 선택한다. R bit는 page가 최근에 참조되었는지, M bit는 수정된 page인지 나타낸다. NUR (0,1 first)는 (0,0) 다음에 (0,1)을 먼저 보고, NUR (1,0 first)는 (0,0) 다음에 (1,0)을 먼저 본다. 이 차이는 최근 사용 여부를 더 중시할 것인지, clean page 제거를 더 중시할 것인지의 차이로 해석할 수 있다.")
    code(doc, "NUR (0,1 first): (0,0) -> (0,1) -> (1,0) -> (1,1)\nNUR (1,0 first): (0,0) -> (1,0) -> (0,1) -> (1,1)\nfor each target class in priority order:\n    scan resident frames\n    choose the first page whose (R,M) matches the class")
    para(doc, "NUR에서 중요한 점은 reset interval이다. 본 구현은 index가 reset interval의 배수에 도달하면 resident page의 R bit를 모두 false로 내린다. 이렇게 하지 않으면 한 번 참조된 page들이 계속 R=1로 남아 class 구분이 어려워진다. 반대로 reset이 너무 자주 일어나면 최근 참조 정보가 너무 빨리 사라진다. 따라서 NUR는 class 우선순위뿐 아니라 reset 주기까지 함께 생각해야 한다.")
    para(doc, "Second Chance는 clockHand가 가리키는 frame부터 page를 검사한다. 검사한 page의 R bit가 1이면 그 bit를 0으로 내리고 다음 frame으로 이동한다. R bit가 0이면 그 page가 victim이 된다. 이 과정은 최근에 참조된 page에게 한 번 더 기회를 주는 방식이다. 하지만 그 기회는 영구적이지 않고, scan 과정에서 R bit가 0으로 내려간 뒤 다음 교체 때는 victim이 될 수 있다.")
    code(doc, "while true:\n    frameData = frame_window[clockHand]\n    if reference bit of frameData is 0:\n        victimIndex = clockHand\n        move clockHand to next frame\n        return victimIndex\n    else:\n        clear reference bit of frameData\n        move clockHand to next frame")
    para(doc, "LRFU-Lite는 page별 score를 유지한다. 모든 resident page의 score는 매 step마다 0.85배로 줄어든다. 현재 참조된 page가 frame 안에 있으면 score가 1.0 증가하고, 새로 들어온 page는 score 1.0으로 시작한다. Victim 선택이 필요하면 score가 가장 낮은 page를 제거한다. 반복 참조된 page는 높은 score를 유지하므로 보호받기 쉽지만, 과거에 많이 쓰인 page가 새 작업 집합보다 오래 보호될 수도 있다.")
    code(doc, "for each resident page:\n    score[page] = score[page] * 0.85\nif current page is hit:\n    score[current page] += 1.0\nif replacement is needed:\n    victim = resident page with the smallest score")

    doc.add_heading("3. 주요 주제의 알고리즘 동작 사례", level=2)
    para(doc, "동작 사례는 Textbook reference string 123412512345, frame size 4를 중심으로 설명할 수 있다. 처음 1, 2, 3, 4가 차례로 들어올 때는 빈 frame이 있으므로 모두 page fault이지만 migration은 아니다. 이후 1과 2가 다시 참조되면 hit가 발생한다. 그러나 FIFO는 hit가 발생해도 순서를 바꾸지 않기 때문에, 7번째 참조 5가 들어올 때 가장 오래된 1을 제거한다.")
    para(doc, "문제는 바로 다음 8번째 참조가 다시 1이라는 점이다. FIFO는 직전에 1을 제거했으므로 다시 page fault가 발생한다. 이 짧은 구간은 FIFO가 최근 사용 정보를 반영하지 못한다는 점을 잘 보여준다. 같은 입력을 NUR로 실행하면 R bit reset 이후의 class를 확인하므로 단순히 오래된 순서만으로 victim을 정하지 않는다. 이 차이는 frame snapshot에서 victim이 바뀌는 지점으로 확인할 수 있다.")
    para(doc, "같은 구간을 NUR 관점에서 보면, page가 frame에 들어오거나 hit될 때 R bit가 1로 설정된다. reset interval이 4이므로 특정 시점마다 R bit가 내려가고, 그 뒤 다시 참조된 page만 R=1이 된다. 이 상태에서 fault가 발생하면 NUR는 page들이 어느 class에 속하는지 확인한다. 그래서 FIFO처럼 무조건 맨 앞 page를 제거하지 않고, 최근 참조되지 않은 page를 먼저 찾게 된다.")
    para(doc, "Second Chance는 이 상황에서 clock hand가 어떤 frame을 먼저 보느냐에 따라 결과가 달라진다. R bit가 1인 page는 한 번 지나가며 살아남지만, 그 과정에서 R bit가 0으로 내려간다. 따라서 현재 step에서 보호받은 page가 다음 replacement에서 다시 보호받는다는 보장은 없다. 이 특징 때문에 clock start 실험에서 fault 수 차이가 나타난다.")
    para(doc, "LRFU-Lite는 Textbook에서는 NUR보다 좋지 않지만, FrequencyBias처럼 A가 반복적으로 등장하는 입력에서는 A의 score를 높게 유지하여 유리하게 작동한다. 반대로 LocalityShift처럼 초반에는 ABC가 자주 나오다가 중간에 DEF로 작업 집합이 바뀌는 입력에서는, 초반에 높아진 score가 남아 있어 새 작업 집합에 적응하는 속도가 늦어진다. 이 사례는 신규 정책도 입력 특성에 따라 장단점이 갈릴 수 있음을 보여준다.")
    para(doc, "이처럼 동작 사례를 보면 page replacement 정책의 차이는 단순히 최종 hit/fault 수에서만 나타나는 것이 아니다. 같은 fault라도 어떤 page가 victim이 되었는지, 그 page가 다음에 다시 등장하는지, 교체 이후 frame 구성이 어떻게 바뀌는지가 중요하다. 따라서 본 프로젝트에서는 step별 로그와 frame snapshot을 함께 출력하도록 하여, 최종 결과의 원인을 따라갈 수 있게 하였다.")

    doc.add_heading("4. 실행 결과 시각화 사례", level=2)
    para(doc, "아래 그림들은 실행파일에서 각 정책과 대표 입력 문자열을 실행한 뒤 Save 기능으로 저장한 시각화 결과를 보고서에 맞게 잘라낸 것이다. 맨 위 행은 reference string의 각 page 참조 순서를 나타내고, 아래 칸들은 각 시점의 frame 상태를 나타낸다. 빨간색은 빈 frame에 새 page가 들어온 page fault, 초록색은 hit, 보라색은 기존 page가 victim으로 선택되어 교체된 migration/replacement를 의미한다.")

    visuals = [
        ("exe_fifo_textbook_crop.png", "동작 그림 1. FIFO 실행 결과: Textbook 입력 123412512345, frame=4", "FIFO에서는 초기 적재 이후 5가 들어오면서 가장 오래된 1이 제거되고, 바로 다음에 1이 다시 참조되어 fault가 발생한다. 이 그림은 FIFO가 최근 hit 정보를 반영하지 못한다는 점을 직관적으로 보여준다."),
        ("exe_nur01_priority_crop.png", "동작 그림 2. NUR (0,1 first) 실행 결과: NURPriority 입력 ABCDABEFABGHABCD, frame=4", "NUR (0,1 first)는 R/M bit class를 기준으로 victim을 찾는다. 이 입력에서는 A와 B가 중간중간 다시 등장하므로, 최근 참조된 page와 modified page가 어떤 class에 남는지에 따라 뒤쪽 hit와 replacement 흐름이 달라진다."),
        ("exe_nur10_clock_crop.png", "동작 그림 3. NUR (1,0 first) 실행 결과: ClockSensitive 입력 ABCDEABCDA, frame=4", "NUR (1,0 first)는 clean page를 상대적으로 먼저 고려하는 변형이다. ClockSensitive 입력에서는 이 우선순위가 유리하게 작용하여 뒤쪽 D와 A 참조가 hit로 이어지는 구간이 나타난다."),
        ("exe_second_clock_crop.png", "동작 그림 4. Second Chance 실행 결과: ClockSensitive 입력 ABCDEABCDA, frame=4, clock start=4", "Second Chance는 clock hand의 시작 위치에 따라 scan 순서가 달라진다. 이 예시는 clock start=4 조건으로, 뒤쪽 A 참조가 hit가 되는 흐름을 보여 주며 clock 기반 정책의 상태 의존성을 설명하기 좋다."),
        ("exe_lrfu_locality_crop.png", "동작 그림 5. LRFU-Lite 실행 결과: LocalityShift 입력 ABCABCABCDEFDEFABC, frame=4", "LRFU-Lite는 초반에 반복된 A, B, C의 score가 높게 남는다. 그래서 중간에 D, E, F로 작업 집합이 바뀌는 구간에서 보라색 교체가 많이 나타나며, 신규 정책의 약점도 함께 확인할 수 있다."),
    ]
    for filename, caption, explanation in visuals:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(ASSETS / filename), width=Inches(6.15))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        set_font(r, size=8.7, italic=True, color="555555")
        para(doc, explanation)


def add_evaluation(doc):
    doc.add_heading("IV. 성능 평가", level=1)
    doc.add_heading("1. 실험 환경", level=2)
    doc.add_heading("1) 실험 도구 및 구현 환경", level=3)
    para(doc, "실험은 제출 패키지에 포함된 C# WinForms 프로젝트와 Release 실행파일을 기준으로 수행하였다. 구현 언어는 C#이며, 주요 알고리즘 로직은 Core.cs에 작성되어 있고, 사용자 입력 처리와 결과 출력은 Form1.cs에 구현되어 있다. Page.cs는 각 참조 단계의 상태를 기록하기 위한 구조체를 제공한다. 실행파일은 Windows 환경에서 동작하며, 사용자는 GUI에서 정책, reference string, frame size, clock start, reset interval, modified pages를 입력한 뒤 Run 버튼으로 실험을 수행한다.")
    para(doc, "실험 도구는 두 가지 방식으로 사용하였다. 첫째, 실행파일을 직접 실행하여 grid 시각화와 chart 결과를 확인하였다. 둘째, 동일한 Core.cs 로직을 사용하여 여러 reference string과 frame size 조합을 반복 실행하고, hit, fault, migration, fault rate, estimated delay 값을 정리하였다. 이렇게 한 이유는 GUI 화면은 동작 과정을 이해하기 좋고, 반복 실행 결과는 표와 그래프로 정리하기 좋기 때문이다.")

    doc.add_heading("2) 실험 제약 사항 및 특징적 고려 사항", level=3)
    para(doc, "본 프로젝트는 실제 운영체제 커널의 page table, TLB, trap 처리, 디스크 I/O를 구현한 것이 아니라 page replacement policy를 관찰하기 위한 시뮬레이터이다. 따라서 page fault가 발생했을 때 실제 디스크에서 page를 읽어 오는 시간이나 dirty page를 write-back하는 시간은 직접 측정하지 않는다. 대신 page fault 1회당 10ms라는 가정값을 사용하여 estimated delay를 계산한다.")
    para(doc, "또한 reference string은 과제 조건에 맞게 한 글자 단위 page로 처리한다. 예를 들어 ABCDAB은 A, B, C, D, A, B 순서로 page를 참조한다. 실제 프로그램의 메모리 주소 trace보다 단순하지만, frame 상태와 victim 선택 과정을 눈으로 추적하기 쉽다는 장점이 있다. Modified Pages도 실제 write instruction을 추적하는 방식이 아니라 사용자가 입력한 page 문자를 modified page로 간주하는 방식이다.")
    para(doc, "본 실험의 특징적 고려 사항은 정책별 고유 입력을 따로 분석했다는 점이다. Second Chance는 clock start 값에 따라 scan 시작 위치가 달라지고, NUR는 R reset interval과 Modified Pages 입력에 따라 R/M class 구성이 달라진다. LRFU-Lite는 별도 UI 입력은 없지만 score decay factor 0.85를 사용하므로, 반복 참조와 작업 집합 전환 입력에서 score가 어떻게 작동하는지 관찰하였다.")

    doc.add_heading("3) 성능 평가 입력 요소와 선정 이유", level=3)
    para(doc, "성능 평가 입력 요소는 reference string, frame size, page replacement policy, clock start, reset interval, modified pages이다. Reference string은 page 참조 순서를 의미하고, frame size는 동시에 메모리에 올릴 수 있는 page 수를 의미한다. Page replacement policy는 victim 선택 기준을 결정한다. Clock start는 Second Chance의 초기 scan 위치이고, reset interval은 NUR에서 R bit를 주기적으로 초기화하는 간격이며, modified pages는 NUR의 M bit를 결정하는 입력이다.")
    para(doc, "Frame size는 3, 4, 5로 설정하였다. Frame이 너무 작으면 모든 정책에서 fault가 많이 발생하고, frame이 충분히 크면 정책 차이가 작아진다. 따라서 3, 4, 5는 작은 시뮬레이터에서 메모리 압박 정도가 달라지는 모습을 보기 위한 범위이다. 기본 조건은 clock start=1, R reset interval=4, Modified Pages=AD로 두었고, 이후 정책 고유 입력의 영향을 보기 위해 해당 값을 따로 변화시켰다.")
    para(doc, "실험 입력은 다섯 가지로 나누었다. Textbook은 강의자료의 FIFO 예시와 연결되는 기본 참조열이다. NURPriority는 NUR class 우선순위 차이를 보기 위한 입력이다. ClockSensitive는 Second Chance의 clock start 영향을 보기 위한 입력이다. FrequencyBias는 A가 반복적으로 등장하여 score 기반 정책의 장점을 보기 위한 입력이다. LocalityShift는 작업 집합이 바뀌는 상황에서 정책이 얼마나 빨리 적응하는지 보기 위한 입력이다.")
    para(doc, "Textbook 입력은 123412512345이다. 이 입력은 짧지만 FIFO의 약점을 보기 좋다. 처음에는 1, 2, 3, 4가 차례대로 들어가고, 이후 1과 2가 다시 참조된다. 그러나 FIFO는 hit가 발생해도 순서를 바꾸지 않으므로, 뒤에서 5가 들어올 때 최근에 다시 사용된 1이 제거될 수 있다. 이 구조 때문에 frame 수를 늘렸는데도 fault가 증가하는 현상이 나타난다.")
    para(doc, "NURPriority 입력은 ABCDABEFABGHABCD이다. 이 입력은 A와 B가 반복적으로 다시 등장하면서도 E, F, G, H 같은 새 page가 들어온다. Modified Pages를 AD로 두면 A와 D는 M=1인 page가 된다. 따라서 NUR가 victim을 선택할 때 단순히 최근 사용 여부만 보는 것이 아니라 modified 여부까지 함께 고려하게 된다.")
    para(doc, "ClockSensitive 입력은 ABCDEABCDA이다. 이 입력은 A부터 E까지 page가 차례로 들어온 뒤 다시 A, B, C, D, A가 등장한다. Frame size가 4일 때 첫 교체 이후 어떤 page가 남는지가 뒤쪽 hit/fault에 큰 영향을 준다. 그래서 Second Chance의 clock hand 시작 위치를 바꾸면 결과가 달라질 수 있다.")
    para(doc, "FrequencyBias 입력은 AAAABCAAADEFAAA이다. A가 매우 많이 등장하므로 반복 참조된 page를 보호하는 정책이 유리할 가능성이 높다. 이 입력은 LRFU-Lite처럼 score를 사용하는 정책이 의도대로 작동하는지 확인하기 위해 넣었다. 동시에 NUR와 Second Chance도 R bit를 통해 A의 최근 참조를 어느 정도 보호할 수 있다.")
    para(doc, "LocalityShift 입력은 ABCABCABCDEFDEFABC이다. 초반에는 ABC가 working set처럼 반복되고, 중간에는 DEF가 등장하며, 마지막에는 다시 ABC로 돌아온다. 이 입력은 과거 반복 정보를 오래 유지하는 정책이 작업 집합 전환 상황에서 불리해질 수 있는지 보기 위해 설계하였다.")

    doc.add_heading("4) 평가 접근 방법과 입출력 요소", level=3)
    para(doc, "평가 접근 방법은 같은 입력을 여러 정책에 적용하여 결과를 비교하고, 이후 특정 정책의 고유 매개변수를 바꾸어 변화 추이를 보는 방식이다. 먼저 다섯 workload와 frame size 3, 4, 5 조합에서 다섯 정책을 모두 실행하였다. 그 다음 Second Chance는 clock start를 1, 2, 3, 4로 바꾸었고, NUR는 reset interval을 2, 3, 4, 6으로 바꾸었다. 마지막으로 Modified Pages를 A, AD, BDF, 없음으로 바꾸어 M bit 설정이 결과에 미치는 영향을 확인하였다.")
    para(doc, "입력 요소는 사용자가 GUI에서 넣는 값과 동일하게 구성하였다. 출력 요소는 hit count, page fault count, migration count, page fault rate, estimated delay, frame snapshot, algorithm state이다. 정량 분석에는 hit, fault, migration, fault rate, delay를 사용했고, 동작 과정 분석에는 frame snapshot과 algorithm state를 사용하였다. 따라서 숫자 결과와 시각화 결과를 함께 해석할 수 있다.")

    doc.add_heading("5) 성능 평가 기준과 선정 이유", level=3)
    para(doc, "평가 기준은 hit count, page fault count, migration count, page fault rate, estimated delay이다. 가장 중요한 기준은 page fault count이다. Page fault가 많을수록 필요한 page를 메모리에서 찾지 못한 경우가 많다는 뜻이고, 가상 메모리 시스템에서 page fault는 일반 메모리 접근보다 훨씬 큰 비용을 발생시키기 때문이다.")
    para(doc, "Page fault rate는 reference string 길이가 다른 입력을 볼 때 비율로 비교하기 위한 지표이다. Migration count는 page fault 중에서도 frame이 가득 차 실제 replacement가 발생한 횟수를 의미한다. 초기 적재에서 발생한 fault와 실제 victim 선택이 필요한 fault를 구분할 수 있기 때문에, 정책의 replacement 부담을 해석하는 데 도움이 된다.")
    para(doc, "Estimated delay는 page fault count에 10ms를 곱한 값이다. 실제 하드웨어 시간 측정값은 아니지만, fault가 늘어날수록 지연 비용도 커진다는 관계를 직관적으로 보여준다. Execution time은 WinForms UI 갱신, chart 생성, 문자열 출력 시간이 섞일 수 있으므로 알고리즘 성능 비교의 핵심 지표로 사용하지 않았다.")

    doc.add_heading("2. 실험 결과 및 분석", level=2)
    doc.add_heading("1) 구체적인 실험 결과", level=3)
    rows = []
    for workload, _, _ in WORKLOADS:
        for row in [r for r in RESULTS if r["workload"] == workload and r["frame"] == 4]:
            rows.append((row["workload"], row["policy"], row["hit"], row["fault"], row["migration"], f'{row["fault_rate"]:.2f}%', row["delay"]))
    metric_table(doc, ["Workload", "Policy", "Hit", "Fault", "Migration", "Fault Rate", "Delay(ms)"], rows, font_size=7.5, left_cols={0, 1})
    para(doc, "Frame=4 기준 결과는 단순히 정책별 순위를 매기기 위한 표가 아니라, 각 입력이 어떤 알고리즘 특징을 드러내는지 보기 위한 기준이다. Textbook은 삽입 순서 기반 정책의 특성을, NURPriority는 R/M bit class 우선순위의 차이를, ClockSensitive는 clock hand 시작 위치와 scan 순서의 영향을, FrequencyBias는 반복 참조 page 보호 효과를, LocalityShift는 작업 집합 변화에 대한 적응성을 보기 위해 사용하였다.")
    para(doc, "Textbook에서는 frame 안에 들어온 순서와 실제 재참조 순서가 엇갈리는 구간이 있어, 삽입 순서만 보는 정책의 동작을 관찰하기 좋다. NUR 두 변형은 이 입력에서 모두 fault 6회로 같은 결과를 보였는데, 이는 (0,1)과 (1,0)의 우선순위 차이보다 R bit를 통해 최근 참조 page를 구분하는 효과가 더 크게 작용했기 때문이다. 즉, NUR의 핵심 특징은 단순 순서가 아니라 bit 상태를 기준으로 victim 후보를 나눈다는 데 있다.")
    add_graph_pair(doc, "aligned_Textbook.png", "그림 1. Textbook frame=4", "aligned_NURPriority.png", "그림 2. NURPriority frame=4")
    para(doc, "NURPriority에서는 NUR (0,1 first)가 fault 10회, NUR (1,0 first)가 fault 11회를 기록하였다. 두 변형의 차이는 1회이지만, 이 차이는 (0,1)과 (1,0)의 우선순위가 실제 결과에 영향을 줄 수 있음을 보여준다. Modified Pages=AD 조건에서 어떤 page가 M=1이 되는지에 따라 class 구성이 달라지고, 그 결과 victim 선택도 달라진다.")
    para(doc, "이 입력에서 중요한 것은 특정 정책이 다른 정책보다 몇 회 더 좋다는 사실보다, NUR가 제공하는 고유 입력인 Modified Pages와 R reset interval이 실제로 결과를 바꾼다는 점이다. A와 D를 modified page로 둔 상태에서는 clean page를 상대적으로 먼저 보는 변형이 항상 유리하지 않았다. 따라서 NUR 분석에서는 fault count뿐 아니라 선택된 victim이 어떤 R/M class였는지 함께 보아야 한다.")
    add_graph_pair(doc, "aligned_ClockSensitive.png", "그림 3. ClockSensitive frame=4", "aligned_FrequencyBias.png", "그림 4. FrequencyBias frame=4")
    para(doc, "ClockSensitive에서는 NUR (1,0 first)가 fault 7회로 가장 낮고, Second Chance는 기본 clock start=1 조건에서 fault 9회를 보였다. 여기서 Second Chance를 해석할 때는 전체 resident page를 한 번에 비교하는 정책이 아니라, clock hand가 만나는 순서대로 R bit를 소비하는 정책이라는 점이 중요하다. 그래서 같은 reference string이라도 clock start를 바꾸면 scan 경로와 후속 frame 구성이 달라질 수 있다.")
    para(doc, "FrequencyBias에서는 A가 매우 자주 반복된다. 이 입력은 반복 참조 page를 보호하는 방식이 어떻게 작동하는지 보기 위한 것이다. LRFU-Lite는 A가 참조될 때마다 score가 증가하므로 A를 victim으로 선택할 가능성이 낮아진다. NUR와 Second Chance도 R bit를 통해 A의 최근 참조 흔적을 남긴다. 다만 입력이 짧고 A 반복이 뚜렷하기 때문에, score 방식의 차이가 매우 크게 벌어지지는 않았다.")
    add_graph_pair(doc, "aligned_LocalityShift.png", "그림 5. LocalityShift frame=4")
    para(doc, "LocalityShift에서는 LRFU-Lite가 fault 12회로 높게 나타난다. 이 입력은 초반 ABC 반복 이후 DEF로 관심 page가 바뀌기 때문에, 과거 반복 정보를 오래 유지하는 정책의 약점을 보여준다. LRFU-Lite의 score는 반복 page 보호에는 도움이 되지만, 작업 집합이 갑자기 바뀌면 이전 page의 높은 score가 새 page의 정착을 방해할 수 있다.")
    para(doc, "따라서 Frame=4 결과를 해석할 때는 모든 정책을 하나의 기준 정책과만 비교하기보다, 각 정책이 가진 고유 상태가 어떤 입력에서 드러나는지 보아야 한다. NUR는 modified pages와 reset interval, Second Chance는 clock start와 scan 경로, LRFU-Lite는 score 누적과 decay, FIFO는 삽입 순서 유지가 핵심 특징이다. 이 특징들이 입력 문자열의 반복성, 지역성 전환, modified page 구성과 만나면서 서로 다른 fault 흐름을 만든다.")

    doc.add_heading("2) Frame 수 변화에 따른 정량 결과 추이", level=3)
    frame_rows = []
    for workload, _, _ in WORKLOADS:
        for policy in POLICIES:
            vals = []
            for frame in FRAMES:
                row = next(r for r in RESULTS if r["workload"] == workload and r["policy"] == policy and r["frame"] == frame)
                vals.append(row["fault"])
            frame_rows.append((workload, policy, *vals))
    metric_table(doc, ["Workload", "Policy", "Frame=3", "Frame=4", "Frame=5"], frame_rows, font_size=7.4, left_cols={0, 1})
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ASSETS / "frame_fault_trends.png"), width=Inches(6.15))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run("그림 6. Frame size 변화에 따른 workload별 page fault 추이")
    set_font(r, size=8.7, italic=True, color="555555")
    para(doc, "위 그래프는 앞의 frame size 변화 표를 그대로 시각화한 것이다. 표는 정확한 수치를 확인하기 좋고, 그래프는 frame 증가에 따라 fault가 감소하는지, 유지되는지, 또는 특정 구간에서 오히려 증가하는지 흐름을 보기 좋다. 특히 같은 frame 변화라도 workload마다 선의 모양이 다르기 때문에, frame 수 효과는 정책만이 아니라 입력 문자열의 구조와 함께 해석해야 한다.")
    para(doc, "Textbook에서는 frame=5가 되면 모든 정책이 fault 5회로 수렴한다. 주요 page를 담을 공간이 충분해지면서 victim 선택의 영향이 줄어든 것이다. 그러나 frame=3에서 frame=4로 넘어가는 구간에서는 정책별 선이 다르게 움직인다. 이 구간은 frame 수가 작업 집합을 완전히 담기에는 애매한 상태라서, 정책의 내부 기준이 결과에 크게 드러난다.")
    para(doc, "NURPriority에서는 NUR 두 변형의 선이 frame 증가에 따라 완만하게 내려간다. 이 입력에서는 A/B 재참조와 modified page 설정이 함께 작용하므로, frame이 늘어날수록 R/M class 기반 선택이 더 안정적으로 작동한다. 반면 어떤 정책은 frame을 늘려도 fault가 거의 줄지 않는데, 이는 단순 용량 증가보다 victim 선택 방식이 여전히 중요하다는 뜻이다.")
    para(doc, "ClockSensitive에서는 frame=5에서 모든 선이 fault 5회로 모인다. 입력에 등장하는 핵심 page를 충분히 담을 수 있는 순간에는 clock hand나 class 우선순위의 차이가 작아진다. 반대로 frame=3과 frame=4에서는 NUR 변형과 LRFU-Lite, Second Chance의 선이 서로 다르게 움직이므로, 이 구간이 정책별 특징을 관찰하기에 더 적합하다.")
    para(doc, "FrequencyBias는 A가 강하게 반복되는 입력이므로 frame 수 변화보다 반복 page 보호 여부가 더 중요한 모양을 보인다. 여러 정책의 선이 거의 평평하게 유지되는 것은 frame을 하나 늘리는 것보다 A를 resident set 안에 유지하는지가 더 큰 영향을 주기 때문이다. LocalityShift에서는 LRFU-Lite 선이 다른 선보다 높게 유지되는데, 이는 score가 과거 반복 page를 오래 보호하여 작업 집합 전환에 늦게 적응하는 특성과 연결된다.")

    doc.add_heading("3) 가변 인자 변경에 따른 결과 변화", level=3)
    clock_rows = []
    for clock_start in [1, 2, 3, 4]:
        result = simulate("ABCDEABCDA", 4, "Second Chance", clock_start=clock_start)
        clock_rows.append((clock_start, result["hit"], result["fault"], result["migration"], f'{result["fault_rate"]:.2f}%'))
    metric_table(doc, ["Clock Start", "Hit", "Fault", "Migration", "Fault Rate"], clock_rows, font_size=8.2)
    add_graph_pair(doc, "aligned_clock_param.png", "그림 7. Second Chance clock start", "aligned_reset_param.png", "그림 8. NUR reset interval")
    para(doc, "Second Chance의 clock start 실험에서는 start=1과 start=2가 fault 9회, start=3이 8회, start=4가 7회를 기록하였다. 이 결과는 특정 start가 항상 좋다는 뜻은 아니다. 다만 같은 입력에서도 clock hand의 초기 위치가 첫 victim 선택과 이후 frame 구성에 영향을 줄 수 있음을 보여준다.")
    para(doc, "Clock start가 결과에 영향을 주는 이유는 Second Chance가 원형 queue를 순서대로 검사하기 때문이다. 처음 어느 frame에서 시작하느냐에 따라 R bit를 먼저 지우는 page가 달라지고, 첫 victim도 달라질 수 있다. 첫 victim이 달라지면 그 뒤에 어떤 page가 frame에 남아 있는지도 달라진다. 결국 초기 상태의 작은 차이가 후속 hit/fault 차이로 이어진다.")
    para(doc, "NUR reset interval 실험에서는 reset=4일 때 NUR (0,1 first)가 fault 10회로 가장 낮았다. Reset이 너무 자주 일어나면 최근 참조 정보가 빨리 사라지고, 너무 늦게 일어나면 대부분의 page가 R=1로 남아 class 구분력이 약해진다. 따라서 R bit를 사용하는 정책에서는 reset 주기가 성능에 영향을 주는 중요한 매개변수이다.")
    para(doc, "Reset interval이 4일 때만 NUR (0,1 first)가 더 좋아진 것은 이 reference string의 재참조 간격과 reset 주기가 비교적 잘 맞았기 때문이다. 실제로 reset 주기는 입력 특성에 따라 달라질 수 있다. 참조가 빠르게 반복되는 입력에서는 너무 빠른 reset이 손해가 될 수 있고, 참조 패턴이 자주 바뀌는 입력에서는 너무 느린 reset이 손해가 될 수 있다.")

    mod_rows = []
    for modified in ["A", "AD", "BDF", ""]:
        r01 = simulate("ABCDABEFABGHABCD", 4, "NUR (0,1 first)", modified_pages=modified)
        r10 = simulate("ABCDABEFABGHABCD", 4, "NUR (1,0 first)", modified_pages=modified)
        mod_rows.append((modified if modified else "(none)", r01["fault"], r10["fault"]))
    metric_table(doc, ["Modified Pages", "NUR (0,1 first) Fault", "NUR (1,0 first) Fault"], mod_rows, font_size=8.2, left_cols={0})
    add_graph_pair(doc, "aligned_modified_param.png", "그림 9. NUR modified pages")
    para(doc, "Modified Pages 입력은 NUR의 M bit를 직접 바꾼다. Modified Pages가 A 또는 없음일 때는 두 NUR 변형이 모두 fault 10회로 같았다. 그러나 AD와 BDF 조건에서는 NUR (1,0 first)가 fault 11회로 증가했다. Clean page를 먼저 제거하는 전략이 write-back 비용 측면에서는 의미가 있을 수 있지만, 본 실험처럼 fault count를 중심으로 보면 항상 좋은 선택은 아니었다.")
    para(doc, "이 결과는 평가 기준의 중요성을 보여준다. 본 보고서의 주요 지표는 page fault count이므로, clean page를 제거하느냐 dirty page를 제거하느냐의 실제 write-back 비용은 별도 지연으로 반영되지 않는다. 만약 dirty page write-back 비용을 estimated delay에 추가한다면 NUR (1,0 first)의 평가가 달라질 수도 있다. 따라서 현재 결과는 fault 수 관점의 결과로 해석해야 한다.")
    para(doc, "전체적으로 보면 하나의 정책이 모든 입력에서 항상 가장 좋지는 않았다. NUR (0,1 first)는 여러 입력에서 안정적이었고, NUR (1,0 first)는 ClockSensitive에서 좋았다. Second Chance는 구현이 비교적 단순하면서 R bit를 활용할 수 있지만 clock hand 경로에 영향을 받았다. LRFU-Lite는 반복 편향 입력에서 장점이 있었지만, 지역성이 빠르게 바뀌는 입력에서는 불리하였다.")
    para(doc, "따라서 본 실험의 핵심 결론은 특정 정책 하나를 무조건 선택해야 한다는 것이 아니다. 정책이 어떤 정보를 사용하고, 그 정보가 현재 workload와 얼마나 잘 맞는지를 보아야 한다. FIFO는 단순하지만 정보가 부족하고, NUR는 bit 정보를 사용하지만 reset 주기에 민감하며, Second Chance는 구현 부담이 낮지만 clock 경로에 영향을 받는다. LRFU-Lite는 반복성 표현이 가능하지만 과거 정보가 오래 남을 수 있다.")


def add_conclusion(doc):
    doc.add_heading("V. 결론", level=1)
    doc.add_heading("1. 프로젝트 요약", level=2)
    para(doc, "본 프로젝트의 배경은 가상 메모리 시스템에서 page fault가 발생했을 때 어떤 page를 내보낼 것인지 결정하는 page replacement 문제이다. 운영체제는 모든 page를 물리 메모리에 올려 둘 수 없기 때문에 제한된 frame 안에서 참조 가능성이 높은 page를 유지해야 한다. 이때 replacement policy가 부적절하면 같은 프로그램이라도 page fault가 늘어나고, 그 결과 디스크 접근에 해당하는 큰 비용이 발생한다.")
    para(doc, "프로젝트의 목표는 여러 page replacement policy를 하나의 시뮬레이터에서 실행해 보고, 같은 reference string에서도 정책의 내부 기준에 따라 frame 상태와 page fault 흐름이 어떻게 달라지는지 분석하는 것이다. 이를 위해 C# WinForms 기반 프로그램을 사용하였고, GUI에는 reference string, frame size, policy, clock start, reset interval, modified pages 입력을 배치하였다. 실행 후에는 단계별 frame 상태, hit/fault 여부, victim, 정책 상태, hit count, page fault count, migration count, fault rate, estimated delay가 출력되도록 하였다.")
    para(doc, "구현한 알고리즘은 FIFO, NUR (0,1 first), NUR (1,0 first), Second Chance, LRFU-Lite이다. FIFO는 가장 먼저 들어온 page를 제거하는 단순한 기준을 사용한다. NUR는 R bit와 M bit를 이용하여 page를 class로 나누고, 두 변형은 (0,1)과 (1,0) 중 어느 class를 먼저 볼 것인지가 다르다. Second Chance는 FIFO queue 구조를 유지하되 R bit가 1인 page에는 한 번 더 기회를 주고 clock hand를 이동한다. LRFU-Lite는 최근성과 빈도성을 간단한 score로 합쳐 반복적으로 참조된 page를 보호하도록 설계하였다.")
    para(doc, "실험은 Textbook, NURPriority, ClockSensitive, FrequencyBias, LocalityShift 다섯 reference string을 중심으로 수행하였다. 각 입력은 단순히 임의로 만든 문자열이 아니라, 특정 정책의 특징이 드러나도록 설계하였다. Textbook은 삽입 순서 기반 정책의 약점을 보기 위한 입력이고, NURPriority는 R/M class 우선순위와 modified page 설정을 보기 위한 입력이다. ClockSensitive는 clock hand 시작 위치의 영향을, FrequencyBias는 반복 page 보호 효과를, LocalityShift는 작업 집합 전환 상황에서의 적응성을 보기 위한 입력이다.")

    doc.add_heading("2. 핵심 결론", level=2)
    para(doc, "첫 번째 결론은 page fault 수가 정책 이름만으로 결정되지 않는다는 점이다. 같은 정책이라도 reference string의 반복 구조, frame size, reset interval, clock start, modified pages에 따라 결과가 달라졌다. 특히 frame size가 커지면 일반적으로 fault가 줄어들 가능성이 높지만, 모든 경우에 단조롭게 감소하지는 않았다. 이는 단순히 frame 수를 늘리는 것보다 그 frame 안에 어떤 page가 남아 있는지가 더 중요할 수 있음을 보여준다.")
    para(doc, "두 번째 결론은 정책별로 잘 드러나는 입력이 다르다는 점이다. FIFO는 구현과 해석이 쉽지만 최근 참조 정보를 사용하지 않기 때문에 재참조 가능성이 높은 page도 오래전에 들어왔다는 이유만으로 제거할 수 있다. NUR는 R/M bit의 조합을 class로 나누어 victim 후보를 고르지만, reset interval과 modified pages 설정에 민감하다. Second Chance는 원형 queue와 R bit를 함께 사용하므로 clock hand가 어느 위치에서 시작하는지에 따라 scan 경로가 달라질 수 있다. LRFU-Lite는 반복 page를 보호하는 데 유리하지만, 입력의 지역성이 빠르게 바뀌면 과거 score가 새 작업 집합으로의 전환을 방해할 수 있다.")
    para(doc, "세 번째 결론은 정량 결과와 동작 과정 시각화를 함께 보아야 한다는 점이다. 표와 그래프는 fault 수의 차이를 한눈에 보여 주지만, 왜 그 결과가 나왔는지는 단계별 frame snapshot과 algorithm state를 보아야 설명할 수 있다. 예를 들어 NUR의 결과를 해석하려면 victim이 선택된 시점의 R/M class를 보아야 하고, Second Chance의 결과를 해석하려면 clock hand가 어떤 순서로 page를 검사했는지 보아야 한다. LRFU-Lite의 경우에는 score가 어떤 page에 누적되어 있는지 확인해야 한다.")
    para(doc, "네 번째 결론은 본 시뮬레이터의 estimated delay가 실제 시간 측정보다는 page fault 비용을 직관적으로 표현하기 위한 지표라는 점이다. 본 프로그램에서는 page fault 1회당 10ms를 곱하여 delay를 계산한다. 따라서 delay는 fault count와 같은 방향으로 움직이며, 실제 디스크 장치의 지연 시간이나 dirty page write-back 비용을 직접 반영하지는 않는다. 그럼에도 page fault가 증가하면 시스템 비용도 증가한다는 관계를 설명하는 데에는 충분히 의미가 있다.")
    para(doc, "종합하면, 본 프로젝트에서 얻은 핵심 결론은 하나의 replacement policy가 모든 입력에서 항상 우수하지 않다는 것이다. 정책은 각자 사용하는 정보가 다르고, 그 정보가 workload의 특성과 맞을 때 좋은 결과를 낸다. 따라서 페이지 교체 정책을 평가할 때는 단순히 평균 fault 수만 보는 것보다, 입력의 지역성, 반복성, modified page 구성, 정책별 상태 변화까지 함께 분석해야 한다.")

    doc.add_heading("3. 향후 개선 방안 및 추가 실험", level=2)
    para(doc, "향후 개선 방향으로는 먼저 LRFU-Lite의 decay factor를 사용자 입력으로 제공하는 방법이 있다. 현재 구현은 decay factor를 0.85로 고정하여 score를 계산한다. 이 값은 최근성과 빈도성 사이의 균형을 결정하므로, workload에 따라 적절한 값이 달라질 수 있다. 사용자가 decay factor를 바꾸어 실험할 수 있다면 FrequencyBias처럼 반복 참조가 강한 입력과 LocalityShift처럼 작업 집합이 바뀌는 입력에서 어떤 값이 적절한지 더 구체적으로 분석할 수 있다.")
    para(doc, "두 번째 개선 방향은 Modified Pages 입력을 read/write trace로 확장하는 것이다. 현재 프로그램은 사용자가 입력한 page 문자를 modified page로 간주한다. 이 방식은 NUR의 M bit 개념을 설명하기에는 간단하고 직관적이지만, 실제 운영체제에서 dirty bit가 write 접근에 의해 설정되는 방식과는 차이가 있다. Reference string을 단순 page 문자뿐 아니라 read/write 정보와 함께 받으면, NUR의 modified bit와 dirty page 교체 비용을 더 현실적으로 모델링할 수 있다.")
    para(doc, "세 번째 개선 방향은 estimated delay 계산식을 확장하는 것이다. 현재 delay는 page fault count에 10ms를 곱한 값이다. 그러나 실제 시스템에서는 clean page를 제거할 때와 dirty page를 제거할 때 비용이 다를 수 있다. Dirty page는 보조기억장치에 다시 기록해야 하므로 write-back 비용이 추가된다. 이 비용을 별도 항으로 넣으면 NUR (0,1 first)와 NUR (1,0 first)의 차이를 fault count뿐 아니라 비용 관점에서도 비교할 수 있다.")
    para(doc, "추가 실험으로는 더 긴 reference trace를 사용하는 방법이 있다. 본 보고서의 입력은 알고리즘 동작을 눈으로 추적하기 위해 비교적 짧게 구성하였다. 짧은 입력은 설명에는 좋지만 실제 프로그램의 긴 메모리 접근 패턴을 모두 대표하지는 못한다. CSV 파일이나 텍스트 파일에서 긴 trace를 읽어 오고, 결과를 다시 CSV로 저장하면 많은 입력에 대해 평균 fault rate와 분산을 계산할 수 있다.")
    para(doc, "마지막으로 UI 출력과 알고리즘 실행 시간을 분리하는 개선도 가능하다. 현재 execution time에는 WinForms 화면 갱신, DataGridView 출력, chart 생성, 문자열 처리 시간이 함께 섞일 수 있다. 알고리즘 자체의 계산 비용을 비교하려면 Core.cs의 policy 실행만 반복하는 별도 benchmark 모드가 필요하다. 다만 이번 과제의 중심은 실제 시간 최적화보다 replacement 원리와 결과 해석이므로, 현재 보고서에서는 page fault count, migration count, fault rate, estimated delay를 중심으로 결론을 도출하였다.")

    doc.add_heading("참고 자료", level=1)
    for ref in [
        "Ch3. Memory Management and Virtual Memory.pdf, 운영체제 강의자료.",
        "Term Project - Page Replacement Policy Design.pdf, 운영체제 Term Project 과제 안내.",
        "프로젝트 보고서 양식.pdf, 프로젝트 보고서 작성 양식.",
        "제출 소스 파일: Core.cs, Form1.cs, Form1.Designer.cs, Page.cs.",
    ]:
        bullet(doc, ref)


def build():
    doc = Document()
    setup_doc(doc)
    cover(doc)
    add_summary(doc)
    add_intro(doc)
    add_background(doc)
    add_main(doc)
    add_evaluation(doc)
    add_conclusion(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
